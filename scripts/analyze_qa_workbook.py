import json
import sys

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).split())


def main(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    summary = {"Pass": 0, "Fail": 0, "Blocked": 0, "Skip": 0}
    exceptions = []
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or "Status" not in header:
            continue
        index = {clean(value): position for position, value in enumerate(header)}
        for row in rows:
            status = clean(row[index["Status"]])
            if status in summary:
                summary[status] += 1
            if status in {"Fail", "Blocked", "Skip"}:
                exceptions.append({
                    "sheet": worksheet.title,
                    "test_id": clean(row[index.get("Test ID", 0)]),
                    "status": status,
                    "test": clean(row[index.get("Test & Preconditions", 3)]),
                    "expected": clean(row[index.get("Expected Result", 5)]),
                    "actual": clean(row[index.get("Actual Result", 6)]),
                    "notes": clean(row[index.get("Notes / Defect ID", 9)]),
                })
    print(json.dumps({"sheets": workbook.sheetnames, "summary": summary, "exceptions": exceptions}, indent=2))


if __name__ == "__main__":
    main(sys.argv[1])
