from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts" / "qa" / "VoiceVault_Master_QA_Test_Plan_v1.0.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
GOLD = "D99A1B"
DARK = "17212B"
WHITE = "FFFFFF"
GREY = "5B6570"
PALE = "F7F9FB"
RED = "C62828"
GREEN = "2E7D32"


@dataclass(frozen=True)
class TestCase:
    area: str
    title: str
    priority: str
    test_type: str
    preconditions: str
    steps: str
    expected: str
    automation: str = "Candidate"
    requirement: str = ""


CASES: list[TestCase] = []


def add(
    area: str,
    title: str,
    priority: str,
    test_type: str,
    preconditions: str,
    steps: str,
    expected: str,
    automation: str = "Candidate",
    requirement: str = "",
) -> None:
    CASES.append(
        TestCase(area, title, priority, test_type, preconditions, steps, expected, automation, requirement)
    )


def add_route_access_cases() -> None:
    public_routes = [
        ("/", "landing page"),
        ("/pricing", "pricing page"),
        ("/privacy", "privacy notice"),
        ("/terms", "terms of service"),
        ("/login", "login page"),
        ("/signup", "sign-up page"),
        ("/forgot-password", "forgot-password page"),
    ]
    for route, label in public_routes:
        add("PUB", f"Open public {label}", "P1", "Functional", "No authenticated session.",
            f"1. Open {route} directly in a new browser session. 2. Refresh the page. 3. Inspect navigation, primary content, footer, and browser console.",
            "The route returns a successful page, renders its complete content without hydration/runtime errors, and remains available after refresh.", "Automate", f"FE:{route}")
    protected = [
        ("/dashboard", "customer dashboard"), ("/record", "guided recording"),
        ("/processing", "AI processing"), ("/chat", "AI chat"),
        ("/family", "family management"), ("/settings", "account settings"),
    ]
    for route, label in protected:
        add("NAV", f"Block anonymous access to {label}", "P0", "Security",
            "No access or refresh token is stored.",
            f"1. Enter {route} directly. 2. Observe the redirect target. 3. Sign in. 4. Verify whether the intended destination is restored.",
            "Protected data is never rendered to the anonymous user; the user is sent to login and, where supported, safely returned to the intended route after authentication.", "Automate", f"AUTHZ:{route}")
    admin_routes = ["/admin", "/admin/users", "/admin/questions", "/admin/processing", "/admin/payments", "/admin/logs"]
    for route in admin_routes:
        add("ADM-AUTH", f"Enforce administrator access on {route}", "P0", "Authorization",
            "Prepare anonymous, ordinary customer, and allow-listed administrator accounts.",
            f"1. Open {route} as each role. 2. Repeat through client-side navigation and a direct refresh. 3. Call the backing admin API using each role's token.",
            "Anonymous and customer identities receive a safe redirect/403 without admin data leakage; only the allow-listed administrator can view the page and API response.", "Automate", f"ADMIN:{route}")


def add_auth_cases() -> None:
    add("AUTH", "Register a valid new customer", "P0", "Functional", "Use an unused, deliverable email address.",
        "1. Open /signup. 2. Enter full name, valid email, optional phone, compliant password and matching confirmation. 3. Accept terms. 4. Submit once.",
        "One account is created, the response contains the expected session/user payload, protected navigation succeeds, default Free entitlements are present, and no duplicate submission occurs.", "Automate", "POST /api/users/signup/")
    invalid_signup = [
        ("blank full name", "Leave full name empty", "Full name is required; no request creates a user."),
        ("malformed email", "Enter user-at-example", "Email validation is shown and no account is created."),
        ("blank email", "Leave email empty", "Email is required."),
        ("short password", "Enter fewer than six characters", "The password is rejected by the active client/server minimum."),
        ("mismatched confirmation", "Use two different passwords", "The mismatch is shown and submission is blocked."),
        ("terms not accepted", "Leave the terms checkbox clear", "Consent is required and registration is blocked."),
        ("duplicate email with case variation", "Use an existing address with different letter casing", "A second identity is not created and a non-enumerating conflict message is returned."),
        ("leading/trailing whitespace", "Pad name and email with spaces", "Safe fields are trimmed consistently and the canonical email is stored."),
        ("very long values", "Enter name/email/phone at and beyond supported lengths", "Boundary values are handled without truncation ambiguity or server error."),
        ("HTML/script input", "Enter markup in full name and phone", "Input is encoded or rejected; no script executes in any later view."),
    ]
    for title, action, expected in invalid_signup:
        add("AUTH", f"Sign-up validation: {title}", "P1" if "script" not in title else "P0", "Validation",
            "Sign-up page is available; no account should be created by the test.",
            f"1. Populate every other field validly. 2. {action}. 3. Submit by button and Enter key. 4. Check the network response and database/account list.",
            expected, "Automate", "FE:/signup + POST /api/users/signup/")
    add("AUTH", "Password strength meter feedback", "P2", "Usability", "Open /signup.",
        "1. Enter passwords that add length, uppercase, lowercase, digit, and special character one criterion at a time. 2. Remove each criterion. 3. Paste a password.",
        "The meter updates immediately and consistently, accurately explains unmet criteria, remains accessible to assistive technology, and does not expose the password value.", "Automate", "PasswordStrengthMeter")
    add("AUTH", "Log in with valid credentials", "P0", "Functional", "An active customer account exists.",
        "1. Open /login. 2. Enter the registered email and correct password. 3. Submit. 4. Refresh the destination page.",
        "Login succeeds once, access/refresh session state is stored as designed, the current profile loads, and refresh preserves access without exposing tokens in the URL.", "Automate", "POST /api/users/login/")
    login_failures = [
        ("unknown email", "an unregistered email and plausible password"),
        ("wrong password", "a registered email and incorrect password"),
        ("malformed email", "an invalid email shape"),
        ("blank password", "a valid email and blank password"),
        ("inactive/deleted account", "credentials for an inactive or removed account"),
        ("SQL injection payload", "SQL metacharacters in both fields"),
        ("oversized credentials", "values far beyond field limits"),
    ]
    for title, data in login_failures:
        add("AUTH", f"Reject login with {title}", "P0" if title in {"SQL injection payload", "inactive/deleted account"} else "P1", "Negative/Security",
            "No valid session exists.",
            f"1. Submit {data}. 2. Compare UI and API response. 3. Refresh and open /dashboard. 4. Check that logs redact credential values.",
            "Authentication fails without creating session state; the message is useful but does not reveal whether an account exists; secrets and raw credentials are absent from logs.", "Automate", "POST /api/users/login/")
    add("AUTH", "Log out from an active session", "P0", "Functional", "Customer is logged in on a protected page.",
        "1. Trigger logout. 2. Use browser Back. 3. Refresh a previously protected URL. 4. Reuse the old access token against /profile. 5. Attempt refresh with the prior refresh token.",
        "Client state is cleared, protected content cannot be recovered from history/cache, and server-side token behavior matches the documented invalidation policy.", "Automate", "POST /api/users/logout/")
    add("AUTH", "Refresh an unexpired session token", "P0", "API", "Valid refresh token; access token is near expiry.",
        "1. Make an authenticated request after the access token expires. 2. Observe one refresh request. 3. Let the original request retry. 4. issue two parallel API calls.",
        "A valid replacement access token is obtained, the original request succeeds once, parallel calls do not create a refresh storm, and the UI remains signed in.", "Automate", "POST /api/users/refresh/")
    for condition in ["expired refresh token", "malformed token", "access token supplied as refresh token", "token signed with the wrong key", "token missing subject", "revoked/deleted user"]:
        add("AUTH", f"Reject refresh: {condition}", "P0", "API Security", f"Prepare a {condition}.",
            "1. POST the token to /api/users/refresh/. 2. Retry a protected request. 3. Inspect response body, browser state, and logs.",
            "The request is rejected with a stable 4xx response, no new token is issued, local session state is cleared when appropriate, and diagnostic output contains no token material.", "Automate", "POST /api/users/refresh/")
    add("AUTH", "Forgot-password request for known and unknown email", "P1", "Functional/Security", "Prepare one registered address and one unregistered address.",
        "1. Submit each address from /forgot-password. 2. Compare response status, wording, and response time. 3. Check outbound mail only for the registered account. 4. Repeat rapidly.",
        "The UI does not enumerate accounts, a reset message is sent only when appropriate, rate limits apply, and repeated requests do not create unlimited valid reset credentials.", "Candidate", "FE:/forgot-password")
    add("AUTH", "Thirty-minute idle-session timeout", "P0", "Security", "Log in and keep the tab idle; control system clock in test environment.",
        "1. Stay inactive just below 30 minutes and perform an action. 2. Repeat with inactivity beyond 30 minutes. 3. Test in two tabs. 4. Resume from sleep.",
        "Activity before the boundary keeps the session usable; inactivity beyond the boundary clears the session in all tabs and requires authentication before protected data is shown.", "Automate", "voicevault_session idle timeout")
    add("AUTH", "JWT expiry boundary and 60-second buffer", "P0", "API Security", "Create tokens expiring in 61, 60, 59, and 0 seconds.",
        "1. Load the app with each token. 2. Request /profile. 3. Observe refresh/redirect behavior. 4. Verify clock-skew handling.",
        "Tokens inside the configured safety buffer are refreshed or rejected predictably; expired credentials never authorize a request; no looping navigation occurs.", "Automate", "JWT middleware")


def add_consent_profile_cases() -> None:
    consent_types = ["voice_cloning", "ai_personality_generation", "family_access", "terms", "privacy"]
    for consent in consent_types:
        add("CNS", f"Record {consent} consent", "P0", "Privacy", "Authenticated customer; consent is not yet recorded.",
            f"1. Trigger the UI control for {consent}. 2. Read the disclosed purpose. 3. Accept once. 4. Query the profile/consent record. 5. Repeat acceptance.",
            "One auditable consent record is associated with the correct user and type, includes version/time plus permitted IP/user-agent metadata, and duplicate actions do not corrupt state.", "Automate", "POST /api/users/consent/")
        add("CNS", f"Do not infer {consent} consent", "P0", "Privacy", "Authenticated customer without this consent.",
            f"1. Navigate away from the {consent} prompt without accepting. 2. Refresh. 3. Attempt the protected downstream action directly through UI and API.",
            "Consent remains false, the protected action is blocked with an actionable explanation, and browsing/dismissal is not stored as acceptance.", "Automate", "ConsentRecord")
    add("PRO", "Load authenticated profile", "P0", "API", "Authenticated Free customer with known profile and quota values.",
        "1. GET /api/users/profile/. 2. Compare fields to persisted values. 3. Inspect browser network payload. 4. Repeat as another user.",
        "Only the caller's supported profile, plan, processing, payment-summary, and quota fields are returned; password hashes, service keys, and other users' data are absent.", "Automate", "GET /api/users/profile/")
    profile_fields = [("full name", "a normal Unicode personal name"), ("phone number", "a valid international number"), ("full name boundaries", "blank, whitespace-only, maximum, and over-maximum names"), ("phone boundaries", "valid, blank, malformed, and overlong numbers"), ("unsafe markup", "HTML and script-like text")]
    for label, data in profile_fields:
        add("PRO", f"Update profile: {label}", "P1", "Validation", "Authenticated customer.",
            f"1. PATCH /api/users/profile/ with {data}. 2. Reload the UI and API. 3. Log out/in. 4. View the value from admin where permitted.",
            "Valid data is normalized and persists consistently; invalid/unsafe data receives field-level 4xx feedback and is never rendered as executable markup.", "Automate", "PATCH /api/users/profile/")


def add_plan_quota_cases() -> None:
    limits = [
        ("recording questions", 5, 30), ("recording minutes", 15, 300),
        ("monthly text messages", 5, 1000), ("monthly voice responses", 0, 200),
        ("family members", 1, 10), ("storage MB", 500, 5000), ("AI generations", 1, 3),
    ]
    for metric, free_limit, premium_limit in limits:
        for plan, limit in [("Free", free_limit), ("Premium", premium_limit)]:
            add("QTA", f"{plan} {metric} entitlement boundary", "P0", "Business Rule",
                f"{plan} account with usage controllable at {limit - 1 if limit else 0}, {limit}, and {limit + 1} units.",
                f"1. Set {metric} usage just below the limit and perform the action. 2. Reach exactly {limit}. 3. Attempt one additional unit. 4. Refresh profile/quota UI and retry through the API.",
                f"Usage up to the documented {plan} limit ({limit}) is accounted consistently; the first over-limit action is atomically blocked with upgrade/reset guidance and cannot be bypassed by a direct or concurrent API call.", "Automate", "UsageQuota + services/plan_limits.py")
    add("QTA", "Monthly quota reset", "P0", "Business Rule", "Account has exhausted monthly chat/voice usage; test clock can cross reset date.",
        "1. Confirm action is blocked before reset. 2. Cross the configured reset boundary. 3. Trigger the next metered action. 4. Inspect all counters and next reset date.",
        "Only monthly counters reset once at the correct boundary; lifetime/storage values remain intact; the newly permitted action increments from the new period.", "Automate", "UsageQuota.reset_date")
    add("QTA", "Atomic quota use under concurrency", "P0", "Reliability", "Account has exactly one unit remaining for each metered operation.",
        "1. Submit two simultaneous requests for the same metered action. 2. Repeat across two browser tabs and direct API calls. 3. Inspect result records and counters.",
        "At most one request consumes the final unit; no negative/over-limit counter, duplicate artifact, or partial billing/processing record is produced.", "Automate", "Quota concurrency")


def add_payment_cases() -> None:
    add("PAY", "Display packages and current plan", "P0", "Functional", "Use anonymous, Free, and Premium sessions.",
        "1. Open /pricing for each session. 2. Inspect package names, price/currency, entitlements, and current-plan state. 3. Compare with GET /api/payments/packages/.",
        "UI matches the server package catalog; current plan is clearly identified; client code does not invent a price or entitlement; unavailable purchase actions are disabled.", "Automate", "GET /api/payments/packages/")
    add("PAY", "Create Stripe checkout for Free customer", "P0", "Integration", "Authenticated Free customer; Stripe test mode configured.",
        "1. Select Premium. 2. POST create-checkout once. 3. Follow the returned Stripe URL. 4. Confirm line item, currency, amount, customer email, success URL, and cancel URL.",
        "One valid Checkout Session is created for the authenticated customer and configured price; trusted server values determine amount/product; redirect targets are allow-listed.", "Automate", "POST /api/payments/create-checkout/")
    checkout_failures = [
        ("anonymous caller", "send no token", "401 and no Stripe session"),
        ("already Premium customer", "use an active Premium user", "duplicate purchase is prevented or explicitly confirmed per policy"),
        ("invalid package/price ID", "tamper with the requested package", "4xx and no arbitrary Stripe price purchase"),
        ("client-supplied amount", "alter amount/currency fields", "server ignores/rejects tampered commercial values"),
        ("Stripe timeout", "simulate provider timeout", "retryable error without local success state"),
        ("Stripe authentication error", "configure an invalid secret in isolated environment", "sanitized configuration/provider error"),
        ("double click", "submit checkout twice rapidly", "no confusing duplicate active checkout/payment outcome"),
    ]
    for title, action, result in checkout_failures:
        add("PAY", f"Checkout creation: {title}", "P0", "Negative/Integration", "Use Stripe test doubles or test mode; no real charge.",
            f"1. {action}. 2. Call checkout from UI and API. 3. Inspect Stripe objects and local Payment records. 4. Retry if relevant.",
            f"The system produces {result}; it never grants Premium access, leaks provider secrets, or leaves a successful local payment for an unconfirmed charge.", "Automate", "POST /api/payments/create-checkout/")
    outcomes = [
        ("successful card", "4242 4242 4242 4242", "succeeded", "Premium/lifetime access is granted once"),
        ("declined card", "Stripe decline test card", "failed", "access remains Free and a useful error appears"),
        ("3-D Secure success", "Stripe authentication-required test card and approve", "succeeded", "access is granted after authentication"),
        ("3-D Secure failure", "Stripe authentication-required card and fail/cancel", "failed or incomplete", "access remains Free"),
        ("customer cancels", "use the Stripe Back/cancel action", "no success", "the app returns safely to pricing/checkout"),
    ]
    for title, instrument, state, result in outcomes:
        add("PAY", f"Stripe checkout outcome: {title}", "P0", "E2E Integration", "Free customer; Stripe test mode; valid package.",
            f"1. Start checkout. 2. Complete using {instrument}. 3. Return to the application. 4. Query payment status/profile and refresh. 5. Revisit the success URL.",
            f"Payment resolves to {state}; {result}; local and Stripe identifiers reconcile; refresh/revisit is idempotent and never double-grants benefits.", "Automate", "Stripe Checkout")
    add("PAY", "Confirm checkout by valid session ID", "P0", "API", "Completed Stripe test Checkout Session belongs to the caller.",
        "1. POST confirm-checkout with the session ID. 2. Repeat the request. 3. Call from a different authenticated account. 4. alter one character of the ID.",
        "Owner confirmation succeeds idempotently; reuse by another user and malformed/foreign sessions are rejected; payment and profile states stay consistent.", "Automate", "POST /api/payments/confirm-checkout/")
    add("PAY", "Payment success page session verification", "P0", "Functional", "Prepare valid, missing, malformed, unpaid, and foreign session_id query values.",
        "1. Open /pricing/success for each value. 2. Observe loading, success/error state, confetti, profile refresh, and dashboard redirect. 3. Refresh mid-verification.",
        "Only a verified paid session displays success and upgrades the correct account; all other values fail safely with recovery guidance; refresh causes no duplicate state change.", "Automate", "FE:/pricing/success")
    webhook_events = ["checkout.session.completed", "payment_intent.succeeded", "payment_intent.payment_failed", "charge.refunded", "unknown supported-version event"]
    for event in webhook_events:
        add("PAY", f"Process Stripe webhook {event}", "P0", "API Integration", "Stripe test signing secret and deterministic event fixture are available.",
            f"1. POST a correctly signed {event} fixture. 2. Repeat the exact event. 3. deliver events out of order. 4. Compare Stripe and local Payment/User state.",
            "A valid event receives the expected 2xx, changes only the intended records, is idempotent under replay, and converges safely when delivery order changes.", "Automate", "POST /api/payments/webhook/")
    for fault in ["missing signature", "invalid signature", "valid signature with altered body", "oversized body", "malformed JSON", "event for an unknown customer"]:
        add("PAY", f"Reject webhook with {fault}", "P0", "Security", f"Construct a request with {fault}.",
            "1. POST to the public webhook endpoint. 2. Inspect status/body/logs. 3. Query users and payments. 4. Retry the same payload.",
            "The request is rejected or safely ignored according to Stripe guidance; no entitlement/payment mutation occurs; sensitive body/signature material is redacted from logs.", "Automate", "POST /api/payments/webhook/")
    for status in ["pending", "succeeded", "failed", "refunded"]:
        add("PAY", f"Expose {status} payment status consistently", "P1", "Functional/API", f"Customer has a latest payment in {status} state.",
            "1. GET payment status and billing. 2. Load dashboard/settings/pricing. 3. Log out/in. 4. Compare plan and entitlement flags.",
            f"All surfaces report a consistent {status} record and permitted entitlement state; internal provider payloads and other customers' billing data remain hidden.", "Automate", "GET /api/payments/status/ + billing/")


def add_dashboard_recording_cases() -> None:
    add("DSH", "Load dashboard for a new Free customer", "P0", "Functional", "Newly registered Free customer with no recordings or AI configuration.",
        "1. Open /dashboard. 2. Refresh. 3. inspect plan badge, quota/usage, empty state, calls to action, and navigation. 4. Resize to mobile.",
        "The dashboard shows accurate Free entitlements and a clear start-recording path without stale AI/chat actions, runtime errors, or horizontal overflow.", "Automate", "FE:/dashboard")
    dashboard_states = [
        ("recordings uploaded and processing pending", "processing CTA and current status"),
        ("processing in progress", "live/in-progress state without duplicate pipeline start"),
        ("processing failed", "failure summary and supported retry path"),
        ("AI ready", "chat and family actions for the correct AI"),
        ("Premium account", "Premium badge and Premium quota values"),
        ("quota exhausted", "limit/upgrade guidance"),
    ]
    for state, outcome in dashboard_states:
        add("DSH", f"Dashboard state: {state}", "P1", "Functional", f"Customer fixture represents {state}.",
            "1. Load and refresh /dashboard. 2. Compare cards/actions with profile and processing APIs. 3. use each visible primary CTA. 4. Navigate back.",
            f"Dashboard presents {outcome}; actions lead to the correct route and never expose another user's AI, recording, or usage state.", "Automate", "FE:/dashboard")
    add("DSH", "Dashboard partial API failure", "P1", "Resilience", "Simulate profile, processing, or recordings endpoint failure independently.",
        "1. Load /dashboard for each failed dependency. 2. Observe skeleton/error/retry behavior. 3. Restore the endpoint and retry. 4. inspect console.",
        "The page does not crash or falsely show success; affected information has an actionable error/retry state, unaffected navigation remains usable, and recovery does not require a full logout.", "Automate", "Dashboard data loading")

    add("REC", "Grant microphone permission and record an answer", "P0", "E2E", "Supported browser with a working microphone; at least one active question.",
        "1. Open /record. 2. Accept microphone permission. 3. Start, pause/resume if available, and stop an answer. 4. Play it back. 5. continue to next question.",
        "Permission is requested in context; visible timer/state controls are accurate; a playable non-empty answer is stored against the correct question; controls prevent invalid transitions.", "Automate/Manual", "AudioRecorder")
    permission_states = [
        ("denied once", "deny the prompt", "clear instructions for enabling permission; no endless prompt"),
        ("permanently blocked", "block permission in browser settings", "actionable browser-specific recovery guidance"),
        ("no microphone device", "remove/disable input devices", "supported error without page crash"),
        ("device disconnected mid-recording", "disconnect active microphone", "recording ends safely and partial data is handled explicitly"),
        ("device changes", "switch input device between answers", "subsequent capture uses the selected/available device correctly"),
    ]
    for title, action, expected in permission_states:
        add("REC", f"Microphone state: {title}", "P0" if "disconnected" in title else "P1", "Device/Negative",
            "Customer is authenticated on /record; browser/device configuration can be controlled.",
            f"1. {action}. 2. Attempt to start recording. 3. Navigate away and back. 4. Recover the device/permission and retry.",
            f"The app provides {expected}; no zero-byte answer is marked complete; recovery works without duplicate answer records or lost completed answers.", "Manual", "MediaDevices/MediaRecorder")
    recording_actions = [
        ("start and stop rapidly", "start then stop within one second", "short/empty capture is rejected or clearly retained according to minimum policy"),
        ("double-click record", "double-click Start", "only one recorder/stream starts"),
        ("double-click stop", "double-click Stop", "only one finalized answer is created"),
        ("long answer", "record through and beyond the supported duration boundary", "timer/limit is enforced without browser memory failure"),
        ("silence", "record only silence", "the user can identify/re-record unusable audio and no false transcript success is implied"),
        ("background/lock", "background or lock the device while recording", "state is explicit and captured data is not silently corrupted"),
        ("incoming interruption", "simulate call/audio focus interruption on mobile", "recording pauses/stops safely with recovery guidance"),
        ("browser refresh", "refresh during an active recording", "warning/recovery behavior prevents accidental silent data loss where possible"),
        ("route navigation", "navigate away during an active recording", "unsaved-work warning or deterministic cleanup occurs"),
    ]
    for title, action, expected in recording_actions:
        add("REC", f"Recording control: {title}", "P1", "Edge/Usability", "Authenticated customer with microphone access.",
            f"1. Begin a guided answer. 2. {action}. 3. Inspect timer/buttons/local draft. 4. attempt playback and continue/re-record.",
            f"{expected}; MediaStream tracks are released when appropriate; the UI and persisted draft stay synchronized.", "Candidate", "AudioRecorder/RecordingControls")
    question_states = [
        ("Free account receives five questions", "Free", 5),
        ("Premium account receives thirty questions", "Premium", 30),
        ("inactive questions excluded", "any", None),
        ("order respected", "any", None),
        ("domain filter respected", "any", None),
        ("question text with Unicode", "any", None),
        ("no active questions", "any", 0),
    ]
    for title, plan, count in question_states:
        expectation = f"exactly {count} question slots" if count is not None else "the configured active/order/domain result"
        add("REC", f"Question selection: {title}", "P0" if count in {5, 30} else "P1", "Business Rule",
            f"Prepare {plan} account and controlled RecordingQuestion fixtures.",
            "1. GET the active question list. 2. Open /record. 3. compare IDs/order/domain/text. 4. Refresh and start a new session.",
            f"API and UI expose {expectation}, without duplicates or inactive records; an empty configuration yields an actionable state rather than a broken recorder.", "Automate", "GET /api/recordings/questions/")
    stepper_actions = ["next after recording", "back to a prior answer", "skip where permitted", "re-record an answer", "play/pause an answer", "delete and replace an answer", "restore local draft after refresh", "discard local draft deliberately"]
    for action in stepper_actions:
        add("REC", f"Guided flow: {action}", "P1", "Functional", "At least three active questions and captured answers are available.",
            f"1. Perform {action}. 2. Move forward and backward across questions. 3. inspect progress count and review screen. 4. refresh where relevant.",
            "Question-to-audio association, completion count, navigation state, playback, and local draft remain accurate; no unrelated answer is overwritten or duplicated.", "Automate/Manual", "QuestionStepper + recording store")
    add("REC", "Enforce minimum guided answers before upload", "P0", "Business Rule", "Customer has fewer completed answers than the configured minimum.",
        "1. Attempt Review/Upload from the guided flow. 2. Call upload directly with too few answer parts. 3. meet the minimum and retry.",
        "UI and server reject an insufficient set consistently with clear remaining-answer guidance; once met, review/upload becomes available exactly once.", "Automate", "RecordingReview/UploadStep")
    add("REC", "Review all answers before upload", "P0", "Functional", "Minimum answers recorded.",
        "1. Open review. 2. play each answer in sequence. 3. verify question labels/durations. 4. re-record one answer. 5. return to review.",
        "Every answer maps to its original question and duration; only the selected answer changes after re-recording; upload summary totals remain correct.", "Automate/Manual", "RecordingReview")
    upload_formats = [
        ("valid WebM", "audio/webm with a valid WebM signature", "accepted"),
        ("valid MP3", "audio/mpeg with a valid MP3 signature", "accepted"),
        ("valid WAV", "audio/wav with a valid RIFF/WAVE signature", "accepted"),
        ("renamed executable", "non-audio bytes named .mp3", "rejected"),
        ("MIME/extension mismatch", "valid bytes with conflicting MIME and extension", "rejected or normalized only by documented rules"),
        ("zero-byte file", "empty audio part", "rejected"),
        ("corrupt/truncated audio", "recognized header with corrupt body", "rejected before downstream processing"),
        ("unsupported format", "AAC/OGG if unsupported", "rejected with supported-format guidance"),
    ]
    for title, payload, outcome in upload_formats:
        add("REC", f"Upload format: {title}", "P0" if outcome != "accepted" else "P1", "API/Security",
            "Authenticated account within recording/storage quota.",
            f"1. POST multipart recording data containing {payload}. 2. Repeat via UI and direct API. 3. inspect storage, DB, response, and processing queue.",
            f"The payload is {outcome}; validation uses extension, declared MIME, and magic signature; rejected files create no orphaned object or processing job.", "Automate", "POST /api/recordings/upload/")
    upload_failures = ["request exceeds 100 MB", "storage quota exceeded", "recording-minute quota exceeded", "network disconnect at 10%", "network disconnect at 99%", "server returns 503", "duplicate retry after ambiguous timeout", "expired token during upload", "parallel uploads in two tabs"]
    for fault in upload_failures:
        add("REC", f"Upload resilience: {fault}", "P0", "Reliability", "Completed local answers are ready for upload; fault injection is enabled.",
            f"1. Trigger {fault}. 2. Observe progress/error state. 3. retry once after restoring service/session. 4. inspect recordings, storage objects, usage counters, and jobs.",
            "The user gets accurate, recoverable status; completed local audio is not discarded prematurely; retries are idempotent; no duplicate recording, object, quota charge, or processing job remains.", "Candidate", "UploadProgress + recordings API")
    add("REC", "List only the caller's recordings", "P0", "Authorization", "Two users each own recordings across statuses.",
        "1. GET recordings as each user. 2. compare IDs/metadata. 3. modify an ID/client filter. 4. request anonymously.",
        "Each authenticated user sees only their authorized recordings; anonymous access is denied; storage internals and other owners' metadata are absent.", "Automate", "GET /api/recordings/")
    add("REC", "Delete an owned recording", "P0", "Functional", "Customer owns a recording with stored object/transcript references; define processing-state policy.",
        "1. DELETE the recording. 2. confirm UI removal. 3. query database/storage/quota. 4. repeat DELETE. 5. attempt processing/chat access to the deleted artifact.",
        "Authorized deletion follows the documented cascade/retention policy, releases applicable quota, is idempotent, and leaves no usable orphan or broken UI reference.", "Automate", "DELETE /api/recordings/{id}/")
    add("REC", "Prevent deletion of another user's recording", "P0", "Authorization", "User A owns a recording; User B is authenticated.",
        "1. As User B, DELETE User A's recording ID. 2. try predictable/invalid UUIDs. 3. list User A's recordings and inspect storage.",
        "The operation returns non-enumerating 403/404 behavior, makes no change, and does not disclose owner or recording metadata.", "Automate", "Object-level authorization")


def add_processing_cases() -> None:
    stages = [
        ("transcription", "OpenAI Whisper", "recording audio", "Transcript"),
        ("personality analysis", "configured OpenAI analysis model", "completed transcript", "personality profile"),
        ("voice cloning", "ElevenLabs", "eligible recordings plus voice-cloning consent", "voice ID/configuration"),
        ("finalization", "internal finalizer", "required prior artifacts", "AI-ready user/configuration"),
    ]
    for stage, provider, prerequisite, artifact in stages:
        add("AIP", f"Complete {stage} successfully", "P0", "Integration", f"Authenticated owner has valid {prerequisite}; provider/test double for {provider} succeeds.",
            f"1. Start {stage} from the supported full-pipeline/step action. 2. poll status. 3. wait for Celery completion. 4. inspect queue, usage tracking, audit logs, and resulting {artifact}. 5. refresh /processing.",
            f"State moves pending → processing → completed in valid order; exactly one {artifact} is associated with the correct user; timestamps/provider usage are recorded; the UI advances without manual data repair.", "Automate", f"AI stage:{stage}")
        for failure in ["provider 401/invalid key", "provider 429/rate limit", "provider timeout", "provider 500", "malformed provider response"]:
            add("AIP", f"{stage}: {failure}", "P0", "Failure Injection", f"Valid {prerequisite}; configure {provider} double to return {failure}.",
                f"1. Start {stage}. 2. observe task retries/backoff and status. 3. inspect user-facing error and redacted logs. 4. restore provider and use Retry once.",
                f"The stage never reports false success; bounded retry policy and retryable/permanent classification are honored; prior artifacts are preserved; recovery produces one valid {artifact} without duplicate cost/queue entries.", "Automate", f"AI provider:{provider}")
    add("AIP", "Run the full AI pipeline in required order", "P0", "E2E", "Owner has valid recordings and all required consents; no completed AI configuration.",
        "1. POST full-pipeline. 2. observe stage order and status timeline. 3. attempt a second full-pipeline request concurrently. 4. wait for completion and open chat.",
        "Transcription precedes personality/voice operations and finalization; prerequisites gate each stage; one active pipeline exists; final state is AI ready and chat can select the resulting AI.", "Automate", "POST /api/admin/process/full-pipeline/{id}/")
    add("AIP", "Block personality analysis without transcript", "P0", "Business Rule", "Recording exists but no completed transcript.",
        "1. Trigger personality step through UI and API. 2. inspect queue/state. 3. create transcript and retry.",
        "The premature request is rejected with a stable prerequisite error and creates no personality artifact; after transcription, the step succeeds normally.", "Automate", "Personality prerequisite")
    for consent in ["voice_cloning", "ai_personality_generation"]:
        add("AIP", f"Block pipeline without {consent} consent", "P0", "Privacy", f"Owner has recordings but has not granted {consent}.",
            "1. Trigger the protected step directly via API and UI. 2. inspect provider calls and artifacts. 3. grant the explicit consent and retry.",
            "Before consent, no provider request or protected artifact is created and the UI asks for informed consent; after consent, normal processing is permitted and audited.", "Automate", "Consent gate")
    status_states = ["pending", "queued", "processing", "completed", "failed", "retrying"]
    for state in status_states:
        add("AIP", f"Render processing status {state}", "P1", "Functional", f"Fixture has a pipeline/step in {state} state.",
            "1. GET status and open /processing. 2. compare stage badge, explanation, timestamps, controls, and timeline order. 3. refresh and navigate away/back.",
            f"API and UI consistently represent {state}; controls are enabled only when valid; completed prior stages remain visible and errors never expose secrets/provider raw payloads.", "Automate", "GET /api/admin/process/status/{id}/")
    add("AIP", "Processing polling lifecycle", "P1", "Reliability", "Pipeline changes state while /processing is open.",
        "1. Observe polling cadence from pending to completed. 2. leave tab in background. 3. restore it. 4. navigate away. 5. simulate one polling failure.",
        "Polling is bounded, stops on terminal state/unmount, avoids overlapping requests, recovers from transient failure, and does not overload the API or freeze the page.", "Automate", "Processing status polling")
    add("AIP", "View transcribed text", "P0", "Functional/Security", "Owner has a completed transcript containing punctuation, Unicode, line breaks, and markup-like words.",
        "1. Open View Transcribed Text. 2. compare full text with stored transcript. 3. test long content/scroll. 4. attempt access as another user.",
        "The owner sees faithful, readable, safely encoded text; layout handles long content; another user receives no transcript data.", "Automate", "Transcript viewer")
    add("AIP", "Retry only a failed processing step", "P0", "Functional", "Earlier steps completed; one later step failed.",
        "1. select Retry. 2. observe which task is queued. 3. click twice and refresh. 4. inspect completed artifacts and usage records.",
        "Only the eligible failed step and required dependents rerun; completed prerequisites are reused; duplicate clicks are idempotent; status/error clears only after genuine success.", "Automate", "POST retry step")
    add("AIP", "Celery worker crash during provider call", "P0", "Recovery", "Processing task is running; provider request behavior is observable.",
        "1. terminate the worker after task start. 2. restart worker. 3. observe acknowledgement/redelivery. 4. inspect artifacts, provider usage, and final status.",
        "Task is recovered or marked failed per acknowledgement policy; no false completed state or duplicate irreversible provider action occurs; operator can safely retry.", "Candidate", "Celery recovery")
    add("AIP", "Thirty-minute task time limit", "P0", "Reliability", "Force a processing task to exceed configured soft/hard time limit.",
        "1. run the slow task. 2. observe worker termination/error mapping. 3. query status. 4. restore normal latency and retry.",
        "The task exits within configured bounds, records a retryable failure without wedging the queue, releases resources, and can later complete once.", "Automate", "CELERY_TASK_TIME_LIMIT")
    add("AIP", "Finalize only when required stages are complete", "P0", "Business Rule", "Create combinations of missing, failed, and completed transcript/personality/voice artifacts.",
        "1. POST finalize for every prerequisite combination. 2. inspect user.ai_ready and AIConfiguration. 3. complete all requirements and retry twice.",
        "Finalization rejects incomplete combinations, never exposes an unusable AI, and sets readiness/timestamps exactly once only when all required artifacts meet policy.", "Automate", "POST /api/admin/process/finalize/{id}/")


def add_chat_cases() -> None:
    add("CHT", "List accessible AIs for owner", "P0", "Functional", "Owner has one AI-ready configuration and one incomplete configuration.",
        "1. Open /chat and GET accessible-ais. 2. inspect selector. 3. attempt to force incomplete AI ID. 4. compare after refresh.",
        "Only ready, authorized AIs are selectable; the owner identity/label is accurate; incomplete or foreign AI IDs cannot start a conversation.", "Automate", "GET /api/family/accessible-ais/")
    add("CHT", "Send a text message and stream AI response", "P0", "E2E", "Authenticated user has access to an AI-ready profile and remaining text quota.",
        "1. enter a normal message. 2. send once. 3. observe streaming tokens and controls. 4. wait for completion. 5. reload conversation history.",
        "One user message and one ordered assistant response persist; streaming renders incrementally without duplicate/lost chunks; quota increments once; final content matches stored conversation.", "Automate", "POST/SSE /api/chat/stream/")
    inputs = [
        ("empty/whitespace", "spaces and newlines only", "send remains disabled or receives 4xx"),
        ("maximum length", "exact supported maximum", "message is accepted without truncation ambiguity"),
        ("over maximum", "one character beyond maximum", "field-level rejection without provider call"),
        ("Unicode and emoji", "multilingual text, emoji, combining marks", "content round-trips correctly"),
        ("HTML/script markup", "script tags and event attributes", "content is displayed as inert text"),
        ("prompt-injection text", "instructions to reveal system prompts/secrets", "system secrets and protected context are not disclosed"),
        ("URLs and line breaks", "links plus multiline content", "safe consistent rendering without layout break"),
    ]
    for title, payload, outcome in inputs:
        add("CHT", f"Chat input: {title}", "P0" if "script" in title or "injection" in title else "P1", "Validation/Security",
            "Authorized user with accessible AI and available text quota.",
            f"1. Submit {payload}. 2. observe client validation, network request, streaming renderer, stored history, and quota. 3. reload the conversation.",
            f"{outcome}; no XSS, secret disclosure, malformed history, duplicate quota charge, or unexpected provider request occurs.", "Automate", "ChatInput/stream endpoint")
    stream_faults = ["OpenAI 401", "OpenAI 429", "OpenAI timeout before first token", "connection drops mid-stream", "malformed SSE event", "server closes without completion marker", "user navigates away mid-stream", "two messages sent concurrently"]
    for fault in stream_faults:
        add("CHT", f"Streaming resilience: {fault}", "P0", "Failure Injection", "Accessible AI and remaining quota; provider/network fault injection enabled.",
            f"1. Send a message. 2. inject {fault}. 3. inspect partial UI/history/quota/logs. 4. retry after recovery. 5. reload conversation.",
            "The UI exits the loading state, explains retryability, never presents a partial response as complete, stores/charges according to one documented outcome, redacts secrets, and allows a clean subsequent response.", "Candidate", "SSE chat stream")
    add("CHT", "Enforce monthly text quota", "P0", "Business Rule", "Free account has 4 of 5 text messages used.",
        "1. send the fifth message. 2. send a sixth through UI and direct API. 3. issue two concurrent fifth-message requests in a fresh fixture.",
        "The fifth succeeds and increments once; over-limit requests are atomically blocked before provider cost, with reset/upgrade guidance and consistent quota UI.", "Automate", "monthly_text_messages")
    add("CHT", "Record and transcribe a voice message", "P0", "Integration", "Premium user with microphone access and voice/text quota; accessible AI.",
        "1. record a clear voice prompt. 2. submit for transcription. 3. edit/confirm if supported. 4. send to chat. 5. inspect stored message and quota.",
        "Valid audio becomes accurate text associated with the caller, is safely editable if designed, sends once, and consumes the documented voice/text units without exposing raw storage paths.", "Automate/Manual", "POST /api/chat/transcribe-voice/")
    voice_faults = ["Free plan with zero allowance", "voice quota exhausted", "microphone denied", "empty audio", "unsupported audio", "corrupt audio", "transcription provider timeout", "expired token", "oversized audio"]
    for fault in voice_faults:
        add("CHT", f"Voice input: {fault}", "P0", "Negative", "Prepare the stated account/device/payload condition.",
            f"1. attempt voice transcription under {fault}. 2. inspect provider calls, UI, usage counter, and stored objects. 3. recover and retry where applicable.",
            "The request is blocked or fails safely before unintended cost; no phantom text/message or quota increment is created; the UI provides a valid recovery/upgrade path.", "Automate/Manual", "Voice transcription")
    add("CHT", "Generate and play an AI voice response", "P0", "Integration", "Premium caller with voice allowance; selected AI has valid ElevenLabs voice ID.",
        "1. complete a text response. 2. request/poll audio status. 3. play, pause, seek, replay, and change volume. 4. reload and play again.",
        "Audio progresses pending → ready, streams only to authorized callers, matches the response/voice, behaves in media controls, and consumes voice quota once.", "Automate/Manual", "audio-status + audio stream")
    audio_states = ["pending generation", "ready", "generation failed", "missing audio", "expired URL/credential", "range/seek request", "another user's audio ID", "concurrent play on mobile"]
    for state in audio_states:
        add("CHT", f"AI audio state: {state}", "P0" if "another" in state else "P1", "Functional/Security",
            f"Conversation response audio is in {state} condition.",
            "1. query audio status. 2. use player controls or direct stream URL. 3. refresh and retry. 4. inspect access control, caching, content type, and UI state.",
            f"The UI/API represent {state} truthfully; unauthorized access is denied; media headers and range behavior are valid where applicable; failures do not break the text conversation.", "Candidate", "GET audio-status/audio stream")
    add("CHT", "List conversation history in correct order", "P0", "Functional", "Multiple conversations/messages exist for the authorized AI, including identical timestamps if supported.",
        "1. GET conversations. 2. compare chronological/pagination order. 3. reload and switch AI. 4. inspect empty state for a new AI.",
        "Only authorized history is returned deterministically, message roles/content/rating/audio metadata are correct, switching AI cannot mix histories, and empty state is clear.", "Automate", "GET /api/chat/conversations/")
    for rating in [1, 2, 3, 4, 5]:
        add("CHT", f"Rate a conversation {rating}/5", "P2", "Functional", "Authorized completed conversation without a rating.",
            f"1. submit rating {rating}. 2. change it to another valid value if supported. 3. reload. 4. attempt rating from another user.",
            f"The authorized rating {rating} is validated/persisted according to update policy and rendered consistently; unauthorized modification is rejected.", "Automate", "POST /api/chat/rate/")
    for invalid in [0, 6, -1, "text", None]:
        add("CHT", f"Reject invalid conversation rating {invalid}", "P2", "Validation", "Authorized completed conversation.",
            f"1. POST rating value {invalid!r}. 2. inspect status/body. 3. reload the original conversation.",
            "A stable 4xx validation response is returned and the prior rating/conversation data remains unchanged.", "Automate", "POST /api/chat/rate/")


def add_family_cases() -> None:
    add("FAM", "Invite one valid family member", "P0", "E2E", "Owner is authenticated, has remaining family quota, and invitee email is deliverable and not already linked.",
        "1. open /family. 2. enter invitee name, email, and relationship. 3. submit once. 4. inspect member status, outbound email/link, audit record, and quota.",
        "One pending member and one cryptographically strong, single-use, seven-day invitation are created for the owner; email contains the correct HTTPS application URL and no secret is logged.", "Automate", "POST /api/family/invite/")
    relationships = ["spouse", "child", "parent", "sibling", "friend"]
    for relation in relationships:
        add("FAM", f"Invite relationship {relation}", "P1", "Business Rule", "Owner has invite capacity and an unused email.",
            f"1. invite a member with relationship {relation}. 2. inspect list and invitation details. 3. accept invite. 4. inspect accepted member card.",
            f"The {relation} value is accepted, persisted, displayed consistently, and does not alter the authorization scope beyond configured family access.", "Automate", "FamilyMember.relationship")
    invalid_invites = [
        ("blank name", "omit name"), ("malformed email", "use invalid email"),
        ("owner's own email", "use owner email"), ("duplicate owner/email", "invite an already pending or accepted email"),
        ("unsupported relationship", "send an arbitrary relationship value"),
        ("unsafe markup", "use HTML/script-like name"), ("overlong values", "exceed field maximums"),
        ("quota exhausted", "use owner at family-member limit"),
    ]
    for title, action in invalid_invites:
        add("FAM", f"Reject family invite: {title}", "P0" if title in {"quota exhausted", "duplicate owner/email", "owner's own email", "unsafe markup"} else "P1", "Validation/Security",
            "Authenticated owner; test fixture matches the invalid condition.",
            f"1. {action}. 2. submit via UI and API. 3. inspect FamilyMember/cache/mail/audit/quota. 4. correct the value and retry.",
            "Invalid invite is rejected with specific safe feedback; no pending member, token, email, audit-success, or quota use is created; corrected data can be submitted once.", "Automate", "POST /api/family/invite/")
    token_states = [
        ("valid unused token", "accept succeeds and links/authenticates the intended invitee according to flow"),
        ("expired after seven days", "acceptance is rejected and owner can resend"),
        ("already used token", "replay is rejected without changing the linked member"),
        ("random token", "non-enumerating invalid response"),
        ("one-character-tampered token", "non-enumerating invalid response"),
        ("token for deleted invitation", "rejected without restoring membership"),
        ("token accepted concurrently", "exactly one acceptance wins"),
    ]
    for state, outcome in token_states:
        add("FAM", f"Invitation token: {state}", "P0", "Security/E2E", f"Prepare a {state}.",
            "1. GET invitation details. 2. open /accept-invite/{token}. 3. submit acceptance. 4. repeat/replay in a second browser. 5. inspect cache, member, user linkage, and audit trail.",
            f"{outcome}; token values never appear in application logs/referrers beyond required URL handling; the owner/member relationship remains consistent.", "Automate", "Invitation token lifecycle")
    add("FAM", "Resend a pending invitation", "P1", "Functional", "Owner has a pending, unaccepted invitation.",
        "1. click Resend. 2. compare old/new tokens and expiry. 3. attempt old token. 4. accept new token. 5. resend rapidly.",
        "One new email/token is issued according to cooldown policy, old credential is invalidated where designed, expiry resets correctly, rate limiting prevents abuse, and membership is not duplicated.", "Automate", "POST /api/family/members/{id}/resend/")
    for state in ["pending member", "accepted member", "member with conversation history"]:
        add("FAM", f"Remove {state}", "P0", "Functional/Privacy", f"Owner has a {state}.",
            "1. choose Remove and inspect confirmation. 2. cancel once. 3. confirm once. 4. repeat DELETE. 5. test former member access and inspect retained conversation policy.",
            "Cancel changes nothing; confirmation revokes access immediately and idempotently, releases applicable quota, and retains/deletes personal data exactly according to policy without breaking owner data.", "Automate", "DELETE /api/family/members/{id}/")
    add("FAM", "Prevent cross-owner member management", "P0", "Authorization", "Owner A and Owner B each have family members.",
        "1. As Owner B, delete/resend Owner A's member ID. 2. attempt sequential IDs. 3. list both owners' members.",
        "Object-level authorization rejects all cross-owner operations with non-enumerating behavior and no email, status, relationship, or conversation metadata leak.", "Automate", "Family object authorization")
    member_states = ["no members", "pending only", "accepted only", "mixed pending/accepted", "search match", "search no result"]
    for state in member_states:
        add("FAM", f"Family list state: {state}", "P2", "Functional", f"Owner fixture has {state}.",
            "1. open /family. 2. use All, Accepted, and Pending filters. 3. search by visible name/email. 4. refresh and inspect counts/actions.",
            f"The {state} list, tabs, counts, empty state, and available actions are accurate; filtering/search does not expose members from another owner.", "Automate", "GET /api/family/members/")
    add("FAM", "Family member lists accessible AIs", "P0", "Authorization", "Accepted family account is linked to one owner with AI ready; another owner also has AI ready.",
        "1. sign in as the family account. 2. GET accessible-ais and open /chat. 3. force another owner's AI ID. 4. revoke membership and retry.",
        "Only AI identities granted by active accepted relationships are returned; unrelated/revoked access is denied immediately and conversation history is scoped correctly.", "Automate", "GET /api/family/accessible-ais/")
    add("FAM", "Frontend/backend member-edit contract", "P1", "Contract", "Frontend exposes family member update function; backend deployed URL set is known.",
        "1. inspect whether edit UI is reachable. 2. send PATCH /api/family/members/{id}/ with valid and invalid changes. 3. compare OpenAPI/docs/client expectations.",
        "Client and server agree on whether update is supported; unsupported PATCH is not exposed as a broken control, or a documented authorized endpoint validates and persists the change.", "Automate", "Known contract risk: family PATCH")


def add_settings_cases() -> None:
    add("SET", "Render account settings", "P1", "Functional", "Authenticated Free and Premium fixtures exist.",
        "1. open /settings for each fixture. 2. inspect profile, security, billing/plan, consent, and danger areas. 3. refresh and use keyboard navigation.",
        "Every field/action reflects the current account and plan, sensitive values are masked/omitted, and sections remain usable after refresh and by keyboard.", "Automate", "FE:/settings")
    add("SET", "Change password with valid current password", "P0", "Security", "Authenticated customer knows current password.",
        "1. enter current password, a compliant new password, and matching confirmation. 2. submit. 3. test old and new credentials in separate sessions. 4. inspect other active sessions.",
        "Password changes once using secure hashing; old credentials fail, new credentials work, and existing tokens are retained/revoked consistently with documented session policy.", "Candidate", "Settings password change")
    invalid_passwords = ["wrong current password", "new password under eight UI characters", "confirmation mismatch", "same as current password", "blank fields", "very long password", "Unicode password", "rapid repeated submissions"]
    for invalid in invalid_passwords:
        add("SET", f"Password change: {invalid}", "P0", "Validation/Security", "Authenticated customer; original password is known.",
            f"1. submit a change using {invalid}. 2. inspect UI/API/logs. 3. log out and try the original password. 4. correct input and retry where valid.",
            "Invalid change creates no credential or session side effect and leaks no passwords; original credentials remain usable; feedback identifies the correct field without revealing stored-secret details.", "Automate", "Password settings")
    add("SET", "Plan and billing links", "P1", "Functional", "Use Free, successful Premium, pending, failed, and refunded payment fixtures.",
        "1. open billing/plan section for each state. 2. inspect CTA/status. 3. use available pricing/checkout links. 4. navigate back.",
        "Displayed plan/payment status matches server truth and only valid next actions are offered; no hidden Stripe identifiers or other user's billing data appear.", "Automate", "Settings billing")


def add_admin_cases() -> None:
    add("ADM", "Load admin dashboard statistics", "P0", "Functional", "Allow-listed administrator; controlled user/recording/payment/processing fixtures.",
        "1. GET /api/admin/stats/ and open /admin. 2. reconcile every count/aggregate against fixtures. 3. create a record and refresh. 4. inspect loading and zero states.",
        "Counts and operational indicators are accurate, consistently defined, current after refresh, and reveal no unnecessary personal data in aggregate views.", "Automate", "GET /api/admin/stats/")
    add("ADM", "Admin dashboard dependency failure", "P1", "Resilience", "Simulate stats API timeout/503.",
        "1. open /admin. 2. observe page and retry control. 3. restore API and retry. 4. inspect request fan-out/console.",
        "Failure is visible and recoverable without fabricated zeroes or page crash; retry loads once; navigation to other admin tools remains available.", "Automate", "Admin dashboard")
    filters = ["search by email", "search by name", "all users", "AI ready", "processing", "failed", "Free plan", "Premium plan", "no matching result"]
    for f in filters:
        add("USR", f"Admin user list: {f}", "P1", "Functional", "Admin is signed in; fixtures cover mixed users/states.",
            f"1. open /admin/users. 2. apply {f}. 3. combine with available pagination/sort. 4. refresh and clear filters. 5. compare GET query parameters/results.",
            "Only matching users appear with accurate status/plan indicators; results are deterministic, filters survive or reset predictably, and no duplicate/missing row occurs across pages.", "Automate", "GET /api/admin/users/")
    add("USR", "Open admin user detail", "P0", "Functional/Privacy", "Admin and target user with recordings, processing, quota, family, payment summary.",
        "1. open the target detail. 2. compare supported fields to source records. 3. try invalid/missing user ID. 4. inspect network payload.",
        "Supported operational data belongs to the selected user and is accurate; password hashes, full tokens, API keys, and unnecessary provider payloads are absent; invalid ID returns safe 404.", "Automate", "GET /api/admin/users/{id}/")
    admin_user_updates = ["change supported plan/status field", "toggle supported AI/processing flag", "submit no-op patch", "submit invalid enum", "submit protected field such as password hash", "concurrent edits by two admins"]
    for action in admin_user_updates:
        add("USR", f"Admin user update: {action}", "P0", "Authorization/Validation", "Admin and target user fixture; audit logging enabled.",
            f"1. PATCH user detail to {action}. 2. inspect response/profile/UI. 3. repeat or race the request. 4. inspect AuditLog actor, target, old/new values.",
            "Only explicitly mutable fields change atomically; invalid/protected changes are rejected; plan/quota implications remain consistent; material changes are audit logged without secrets.", "Automate", "PATCH /api/admin/users/{id}/")
    add("USR", "Delete a user as administrator", "P0", "Destructive/Privacy", "Disposable target user owns recordings, AI config, family links, conversations, payments, and storage objects.",
        "1. open delete confirmation and cancel. 2. confirm deletion. 3. inspect all related records/storage according to retention policy. 4. test target tokens/login and repeat DELETE.",
        "Cancel is non-mutating; confirmed deletion/deactivation follows the approved retention/cascade policy, revokes access immediately, is idempotent, and records the administrator action.", "Manual", "DELETE /api/admin/users/{id}/")
    add("USR", "Prevent administrator self-lockout", "P0", "Safety", "Signed-in administrator targets own account for deletion or privilege/active-state removal.",
        "1. attempt each dangerous self-operation. 2. inspect confirmation/server response. 3. refresh current session.",
        "The system blocks unsafe removal of the last/current administrator or requires an explicit safeguarded workflow; it never silently leaves administration inaccessible.", "Candidate", "Admin safety")

    add("QUE", "Create a recording question", "P0", "Functional", "Admin is signed in; known current question order.",
        "1. create a question with valid text, domain, expected duration, order, and active state. 2. refresh admin list. 3. load customer question API.",
        "One question persists with normalized values in correct order; active questions appear in eligible customer flows and audit information is available.", "Automate", "POST questions base")
    invalid_questions = ["blank text", "whitespace-only text", "overlong text", "unsupported domain", "negative duration", "zero duration", "extreme duration", "negative order", "duplicate order", "HTML/script content"]
    for invalid in invalid_questions:
        add("QUE", f"Reject/handle question: {invalid}", "P1" if "script" not in invalid else "P0", "Validation",
            "Admin is signed in; baseline questions exist.",
            f"1. submit create and update requests containing {invalid}. 2. inspect field feedback and persisted rows. 3. view customer flow. 4. correct and retry.",
            "The value is rejected or normalized by an explicit deterministic rule; no corrupt ordering, XSS, partial record, or broken recording screen results.", "Automate", "Question validation")
    add("QUE", "Edit an existing question", "P0", "Functional", "Admin; existing active question not currently being mutated.",
        "1. change text/domain/duration/order/active individually. 2. save. 3. refresh. 4. fetch customer question list. 5. inspect existing recording associations.",
        "Supported changes persist atomically and appear in new sessions; historical recording/question references remain interpretable; update is audited.", "Automate", "PATCH question")
    add("QUE", "Toggle question inactive", "P0", "Business Rule", "Active question exists and may have historical recordings.",
        "1. deactivate it. 2. request active and all question lists. 3. begin a new recording session. 4. inspect historical recording display. 5. reactivate.",
        "Inactive question is excluded only from new active selection, remains visible to authorized admin/history as required, and reactivation restores it in correct order.", "Automate", "Question active flag")
    add("QUE", "Delete a question", "P0", "Destructive", "Prepare one unused and one historically referenced question.",
        "1. cancel deletion. 2. delete unused question. 3. attempt referenced question deletion. 4. inspect order, history, and audit.",
        "Confirmation prevents accidents; unused deletion is clean; referenced deletion is blocked/soft-deleted/cascaded according to explicit policy without corrupting recordings.", "Manual", "DELETE question")
    reorder_modes = ["drag one item upward", "drag one item downward", "keyboard reorder", "move first to last", "move last to first", "rapid consecutive reorders", "two admins reorder concurrently"]
    for mode in reorder_modes:
        add("QUE", f"Question reorder: {mode}", "P1", "Functional/Accessibility", "At least five questions in a known order.",
            f"1. {mode}. 2. save/reload. 3. inspect reorder API payload. 4. fetch customer list. 5. repeat after conflict where relevant.",
            "A single deterministic contiguous order persists and is announced/operable by keyboard where applicable; no duplicate/lost question occurs; conflicts are resolved visibly.", "Automate/Manual", "POST questions/reorder/")
    add("QUE", "Seed default questions idempotently", "P1", "Admin/API", "Run against empty database and then database with seeded/custom questions.",
        "1. POST seed on empty state. 2. inspect domains/order/count. 3. POST again. 4. add custom questions and seed again.",
        "Required defaults are created once in valid order; repeated seeds do not duplicate/overwrite custom data unless explicitly documented; response reports actual changes.", "Automate", "POST questions/seed/")
    add("QUE", "Bulk update selected questions", "P1", "Admin/API", "Admin; multiple selected questions across domains/statuses.",
        "1. bulk change a supported field. 2. include one invalid/foreign ID. 3. inspect atomicity/result counts. 4. refresh customer/admin lists.",
        "Authorization and validation apply to every item; atomic/partial behavior is explicitly reported; no unselected question changes; ordering and audit trail remain consistent.", "Automate", "POST questions/bulk-update/")
    add("QUE", "Bulk delete selected questions", "P0", "Destructive", "Admin; mix unused and referenced selected questions.",
        "1. cancel confirmation. 2. confirm bulk delete. 3. include invalid IDs and duplicate IDs. 4. inspect result, history, order, and audit.",
        "Deletion follows reference policy consistently, reports each outcome without double counting, preserves unselected/history data, and requires explicit confirmation.", "Manual", "POST questions/bulk-delete/")
    add("QUE", "Export questions", "P2", "Functional", "Admin; Unicode, commas, quotes, line breaks, active/inactive questions exist.",
        "1. trigger export. 2. inspect filename/content type/encoding/headers/row count/order. 3. open in a spreadsheet. 4. check formula-injection values.",
        "Export contains authorized complete data in deterministic order and UTF-8-safe form; cells cannot execute spreadsheet formulas from untrusted text; no secrets/internal IDs beyond specification leak.", "Candidate", "GET questions/export/")

    for status in ["all", "pending", "processing", "completed", "failed"]:
        add("AADM", f"Admin processing filter: {status}", "P1", "Functional", f"Processing fixtures include {status} records.",
            f"1. open /admin/processing. 2. select {status}. 3. paginate/search if available. 4. refresh and open one status detail.",
            f"Only {status} jobs (or all jobs for all) appear with correct user/stage/timestamps/error summary; ordering is deterministic and sensitive provider data is redacted.", "Automate", "GET /api/admin/processing/")
    for action in ["transcribe", "personality", "voice clone", "full pipeline", "retry failed"]:
        add("AADM", f"Admin triggers {action}", "P0", "Admin Operation", "Admin; target user has the required prerequisites or a controlled missing prerequisite.",
            f"1. trigger {action} from admin UI. 2. verify confirmation if costly/destructive. 3. double-click. 4. inspect job/status/audit. 5. repeat with missing prerequisite.",
            "One eligible job is queued and attributed to the administrator; duplicate clicks are idempotent; missing prerequisites fail before provider cost with actionable feedback.", "Automate", "Admin processing actions")
    for batch in ["process pending", "retry failed"]:
        add("AADM", f"Admin batch action: {batch}", "P0", "Batch/Reliability", "Mixed eligible/ineligible user processing fixtures; admin signed in.",
            f"1. invoke {batch}. 2. cancel confirmation once. 3. confirm once. 4. invoke again concurrently. 5. inspect selected jobs/result summary/audit.",
            "Cancel is non-mutating; only eligible records queue once; partial failures are itemized; worker load remains bounded and audit/result counts are accurate.", "Candidate", "Admin batch endpoint")
    for pstatus in ["all", "pending", "succeeded", "failed", "refunded"]:
        add("APAY", f"Admin payment list: {pstatus}", "P1", "Functional/Privacy", f"Payment fixtures cover {pstatus}; admin signed in.",
            f"1. open /admin/payments. 2. filter {pstatus}. 3. search user/payment identifier. 4. paginate and open details where linked.",
            "Rows match server records and filter deterministically; amounts/currency/status/timestamps reconcile; card data, secrets, and full webhook payloads are never displayed.", "Automate", "GET /api/admin/payments/")
    log_filters = ["authentication failures", "processing failures", "payment events", "family invitations", "admin mutations", "no-result query"]
    for lf in log_filters:
        add("LOG", f"Admin logs: {lf}", "P1", "Observability/Security", "Generate known audit/application events; admin signed in.",
            f"1. open /admin/logs and filter/search for {lf}. 2. compare time, severity, actor, action, target, and request metadata. 3. inspect redaction. 4. try CSV/browser copy if available.",
            "Expected events are searchable, chronologically stable, and attributable; passwords, JWTs, SMTP/API keys, invitation tokens, card data, and raw audio/transcripts are redacted.", "Candidate", "GET /api/admin/logs/")
    add("LOG", "Audit log append-only behavior", "P0", "Security", "Existing audit records and admin/operator database access in a controlled environment.",
        "1. attempt supported API update/delete of an AuditLog. 2. perform an audited admin action. 3. compare sequence/timestamps before and after.",
        "Application paths cannot mutate/delete historical audit records; new records append with correct actor/action/target and trustworthy server time.", "Automate", "AuditLog")


def add_api_contract_cases() -> None:
    endpoints = [
        ("POST", "/api/users/signup/", "public", "JSON"), ("POST", "/api/users/login/", "public", "JSON"),
        ("POST", "/api/users/logout/", "user", "JSON"), ("POST", "/api/users/refresh/", "refresh", "JSON"),
        ("GET/PATCH", "/api/users/profile/", "user", "JSON"), ("POST", "/api/users/consent/", "user", "JSON"),
        ("GET/POST", "/api/recordings/", "user", "JSON/multipart"), ("DELETE", "/api/recordings/{id}/", "owner", "JSON"),
        ("GET", "/api/recordings/questions/", "user", "JSON"), ("GET", "/api/family/members/", "user", "JSON"),
        ("POST", "/api/family/invite/", "user", "JSON"), ("GET/POST", "/api/family/invitations/{token}/", "token", "JSON"),
        ("DELETE", "/api/family/members/{id}/", "owner", "JSON"), ("POST", "/api/family/members/{id}/resend/", "owner", "JSON"),
        ("GET", "/api/family/accessible-ais/", "user/family", "JSON"), ("POST/SSE", "/api/chat/stream/", "authorized AI", "event-stream"),
        ("POST", "/api/chat/transcribe-voice/", "authorized AI", "multipart/JSON"), ("GET", "/api/chat/conversations/", "authorized AI", "JSON"),
        ("POST", "/api/chat/rate/", "conversation member", "JSON"), ("GET", "/api/chat/audio-status/", "conversation member", "JSON"),
        ("GET", "/api/chat/audio/{id}/", "conversation member", "audio"), ("GET", "/api/payments/packages/", "public", "JSON"),
        ("POST", "/api/payments/create-checkout/", "user", "JSON"), ("POST", "/api/payments/confirm-checkout/", "user", "JSON"),
        ("POST", "/api/payments/webhook/", "Stripe signature", "JSON"), ("GET", "/api/payments/status/", "user", "JSON"),
        ("GET", "/api/payments/billing/", "user", "JSON"), ("GET", "/api/admin/stats/", "admin", "JSON"),
        ("GET", "/api/admin/users/", "admin", "JSON"), ("GET/PATCH/DELETE", "/api/admin/users/{id}/", "admin", "JSON"),
        ("GET", "/api/admin/processing/", "admin", "JSON"), ("GET", "/api/admin/payments/", "admin", "JSON"),
        ("GET", "/api/admin/logs/", "admin", "JSON"),
    ]
    for method, path, role, media in endpoints:
        add("API", f"Contract: {method} {path}", "P0", "API Contract", f"Prepare valid fixture/credentials for required role: {role}.",
            f"1. send a valid {method} request. 2. validate status, Content-Type ({media}), schema, required/null fields, types, and error envelope. 3. omit auth or use wrong role. 4. send wrong method and malformed body. 5. test trailing slash behavior.",
            "Valid request meets the documented stable contract; unauthorized/wrong-role calls fail before data access; malformed/wrong-method calls return consistent 4xx responses; no HTML traceback or secret fields are exposed.", "Automate", f"API:{method} {path}")
    add("API", "Pagination boundaries", "P1", "API Contract", "Seed list endpoints with 0, 1, exactly one page, and multiple pages of records.",
        "1. request default, first, last, beyond-last, zero, negative, nonnumeric, and oversized page/page-size values. 2. combine with filters. 3. mutate data between pages.",
        "Pagination metadata and deterministic order are correct; invalid bounds return safe validation/defaults; oversized requests are capped; records are not duplicated by unstable ordering.", "Automate", "List endpoints")
    add("API", "JSON content negotiation and malformed bodies", "P1", "API Contract", "Select representative public, user, admin endpoints.",
        "1. send missing/wrong Content-Type, invalid UTF-8, truncated JSON, duplicate keys, JSON array instead of object, and unsupported Accept header. 2. inspect responses/logs.",
        "Requests receive consistent 4xx/406/415 handling without stack traces, partial mutations, parser ambiguity, or sensitive request-body logging.", "Automate", "DRF parsers/renderers")
    add("API", "Database outage mapped to retryable 503", "P0", "Reliability", "Isolated environment can interrupt database connectivity.",
        "1. stop/block database. 2. call representative read and write endpoints. 3. inspect status, Retry-After, body, and logs. 4. restore DB and retry the write.",
        "Connection failures return sanitized 503 with Retry-After: 3 as designed, never masquerade as validation/auth failure, and recovery does not duplicate the write.", "Automate", "DB middleware")
    add("API", "Rate/throttle enforcement", "P0", "Security", "Known configured throttles for authentication, invitation, upload, chat, and admin endpoints.",
        "1. send requests below, at, and above each limit from same user/IP. 2. vary tokens/IP headers. 3. wait reset period. 4. inspect Retry-After and provider calls.",
        "Legitimate traffic below limit succeeds; excess receives consistent 429 without downstream cost; untrusted forwarded headers cannot trivially bypass limits; access resumes after correct window.", "Automate", "DRF throttling")


def add_integration_cases() -> None:
    integrations = [
        ("PostgreSQL/Supabase pooler", "connect, authenticate, migrate, transact, and reconnect", "persistent domain data remains consistent"),
        ("Supabase object storage", "upload, read, delete, permission failure, timeout", "audio objects are private/authorized and DB/object state reconciles"),
        ("Redis", "cache invitation tokens, broker tasks, disconnect/reconnect, eviction", "tokens/tasks behave safely without silent acceptance or loss"),
        ("Celery worker", "route default/transcription/voice/analysis queues and restart", "tasks reach the correct worker and terminal state once"),
        ("Celery beat", "start scheduler twice and run due tasks", "scheduled work is not duplicated and failures are observable"),
        ("OpenAI transcription", "valid audio, model selection, timeout, rate limit, malformed response", "transcript/usage/error mapping is accurate"),
        ("OpenAI personality analysis", "long/short/Unicode transcript and model failure", "structured profile validates before persistence"),
        ("OpenAI chat", "streaming success/failure and model configuration", "safe ordered response with usage tracking"),
        ("ElevenLabs voice clone", "eligible audio, consent, invalid key, quota, timeout", "one authorized voice ID or recoverable failure"),
        ("ElevenLabs TTS", "generate, poll, retrieve, fail, retry", "response audio maps to correct conversation and voice"),
        ("Stripe", "checkout, webhook, reconciliation, test/live isolation", "commercial state is verified and idempotent"),
        ("Brevo SMTP", "TLS login on smtp-relay.brevo.com:587, send, bounce/failure", "invitation/reset mail sends or reports retryable failure without blocking core transaction incorrectly"),
        ("Traefik/Let's Encrypt", "host rules, HTTP redirect, certificate issuance/renewal", "both configured hosts route securely to the intended service"),
    ]
    for name, actions, outcome in integrations:
        add("INT", f"Integration contract: {name}", "P0", "Integration", f"Staging uses provider sandbox/test credentials; provider traffic and application records are observable.",
            f"1. Exercise {actions}. 2. compare outbound request headers/body with provider contract. 3. simulate each documented provider error class. 4. retry/reconcile. 5. inspect logs, usage, and cleanup.",
            f"{outcome}; provider credentials and personal data are minimized/redacted; retries are bounded/idempotent; environment uses the intended endpoint/model/mode.", "Candidate", f"Provider:{name}")
    add("INT", "Brevo invitation email content", "P0", "Email/E2E", "Brevo sandbox/test recipient is available; application public URL is HTTPS.",
        "1. send a family invite. 2. inspect From, Reply-To, recipient, subject, text/HTML bodies, branding, token link, and encoding. 3. open link on desktop/mobile. 4. inspect spam/authentication diagnostics when available.",
        "Mail is accepted by Brevo and delivered with the configured verified sender; content is readable, personalized safely, contains one correct HTTPS invitation URL, and does not expose SMTP credentials or internal hostnames.", "Manual", "utils/email_service.py")
    add("INT", "SMTP failure does not create misleading invite success", "P0", "Failure Injection", "Force DNS, TCP, TLS, authentication, 4xx, and 5xx SMTP failures separately.",
        "1. submit invite for each fault. 2. inspect UI, member/token records, retry scheduling, logs, and resend behavior. 3. restore SMTP and retry.",
        "The system follows one documented transactional policy: either rollback invite creation or preserve an explicit unsent/retryable pending state; it never claims delivery falsely or logs password/token.", "Candidate", "Brevo SMTP")
    add("INT", "Configuration selects intended AI models", "P0", "Configuration", "Staging environment exposes configured OPENAI_CHAT_MODEL, OPENAI_ANALYSIS_MODEL, ELEVENLABS_MODEL without revealing keys.",
        "1. trigger chat, personality, transcription, voice clone/TTS. 2. inspect mocked outbound model parameters and usage records. 3. set unsupported model in isolated environment.",
        "Every operation uses its configured supported model; invalid configuration fails early with a sanitized operator-facing diagnostic, not a false user-level processing failure.", "Automate", "Environment model configuration")


def add_security_privacy_cases() -> None:
    attacks = [
        ("JWT algorithm confusion", "forge none/alternate algorithm tokens"),
        ("JWT subject substitution", "replace sub/email claims without a valid signature"),
        ("token in query string", "supply bearer credentials through URL parameters"),
        ("IDOR recordings", "enumerate another user's recording IDs"),
        ("IDOR transcripts", "request another user's transcript/processing status"),
        ("IDOR conversations/audio", "request another relationship's conversation/audio IDs"),
        ("IDOR family members", "delete/resend another owner's member ID"),
        ("admin allowlist bypass", "alter email casing/claims or use non-admin token"),
        ("SQL injection", "send SQL metacharacters through search/filter/body fields"),
        ("stored XSS", "persist markup in names, questions, transcripts, messages"),
        ("reflected XSS", "send markup in query and error-producing inputs"),
        ("CSRF on state changes", "submit authenticated cross-origin form/fetch requests"),
        ("CORS abuse", "use untrusted Origin and credentialed preflight"),
        ("host header poisoning", "send untrusted Host/X-Forwarded-Host"),
        ("HTTP request smuggling indicators", "send conflicting length/transfer headers through proxy test tooling"),
        ("path traversal", "inject traversal sequences into filenames/IDs"),
        ("malicious audio file", "upload polyglot or embedded active content"),
        ("mass assignment", "add is_admin, premium, quota, owner fields to writable requests"),
        ("open redirect", "tamper login/checkout success/cancel/return URL"),
        ("email header injection", "inject CRLF into names/addresses/subject-capable input"),
        ("log injection", "send newline/structured-log control characters"),
    ]
    for title, action in attacks:
        add("SEC", f"Attack resistance: {title}", "P0", "Security", "Use an authorized security-test environment with representative roles and logging enabled.",
            f"1. {action}. 2. target relevant UI and API paths. 3. verify response, data, provider calls, logs, and subsequent page rendering. 4. repeat anonymously and as lower privilege.",
            "Attack is blocked or rendered inert at the correct trust boundary; no privilege/data/state/provider side effect occurs; response is non-enumerating and logs remain structured, redacted, and actionable.", "Automate/Manual", f"OWASP:{title}")
    headers = ["Strict-Transport-Security", "Content-Security-Policy", "X-Content-Type-Options", "Referrer-Policy", "frame-ancestors/X-Frame-Options", "secure cache controls"]
    for header in headers:
        add("SEC", f"Production security header: {header}", "P0" if header in {"Strict-Transport-Security", "Content-Security-Policy"} else "P1", "Security/Deployment",
            "Production-like DEBUG=False deployment through Traefik HTTPS.",
            f"1. request public, authenticated, error, API, and audio responses over HTTPS. 2. inspect {header}. 3. attempt the protected browser behavior. 4. compare HTTP redirect path.",
            f"{header} is present with an effective, compatible production value on applicable responses; it is not duplicated/conflicted by proxy/app layers and does not break required assets or streaming.", "Automate", "HTTP headers")
    add("SEC", "Secrets are absent from repository/build/client bundle", "P0", "Security", "Production build artifacts and repository working tree are available to security QA.",
        "1. scan tracked files, images, source maps, Next.js static bundles, Docker image history/config, and rendered error pages for known secret patterns. 2. inspect browser network configuration.",
        "SECRET_KEY, database password/URL, Supabase service key, OpenAI/ElevenLabs/Stripe/SMTP credentials and webhook secrets are server-only and absent from distributed artifacts/logs.", "Automate", "Secret management")
    add("SEC", "Production DEBUG disabled", "P0", "Configuration", "Production-like deployment.",
        "1. trigger 404, validation error, unhandled exception, database outage, and provider failure. 2. inspect page/API/logging behavior and environment report.",
        "DEBUG is False; clients receive branded/sanitized errors without settings, stack traces, filesystem paths, SQL, environment variables, or secret values; operators retain correlation details in protected logs.", "Automate", "DEBUG=False")
    add("SEC", "TLS and HTTP-to-HTTPS enforcement", "P0", "Transport Security", "Both Hostinger hostnames resolve to Traefik; certificate is issued.",
        "1. request both hosts over HTTP and HTTPS. 2. inspect redirect code/location, certificate SAN/chain/expiry/protocols. 3. test unknown host and direct IP. 4. verify renewal alerting.",
        "Configured hosts redirect to canonical HTTPS without loops; certificate is trusted and covers intended names; unknown host/IP does not route to VoiceVault; weak TLS is disabled and renewal is monitored.", "Automate/Manual", "Traefik/Let's Encrypt")
    privacy_scenarios = [
        ("data minimization", "inspect every profile/admin/API response for unnecessary personal/provider data"),
        ("transcript confidentiality", "access/search/cache transcript content across roles"),
        ("audio confidentiality", "access object and streaming URLs after logout/revocation"),
        ("family access revocation", "remove member during an active chat/audio session"),
        ("consent evidence", "export/reconcile consent type, version, timestamp, actor context"),
        ("user deletion/retention", "delete a fully populated account and evaluate each data class"),
        ("backup retention", "verify deletion/expiry procedures across backup copies by policy evidence"),
        ("provider data sharing", "inspect outbound payloads to OpenAI, ElevenLabs, Stripe, Brevo"),
    ]
    for title, action in privacy_scenarios:
        add("PRV", f"Privacy control: {title}", "P0", "Privacy", "Approved privacy/retention policy and populated representative test account.",
            f"1. {action}. 2. compare behavior with the privacy notice/consent. 3. test owner, family, admin, anonymous roles. 4. capture evidence without including raw sensitive content.",
            "Collection, access, disclosure, retention, and deletion match the documented lawful purpose and role; revoked/deleted access is enforced promptly; evidence is auditable and safely redacted.", "Manual", f"Privacy:{title}")


def add_nonfunctional_cases() -> None:
    accessibility = [
        ("keyboard-only navigation", "Tab/Shift+Tab/Enter/Space/Escape/arrow keys across every page, dialog, menu, recorder, player, table and drag-reorder control"),
        ("visible focus", "inspect focus indicator against each dark/light/button/error background"),
        ("screen reader landmarks/headings", "navigate landmarks and headings with VoiceOver/NVDA"),
        ("form labels and errors", "submit invalid login/signup/profile/invite/question inputs and inspect name/description/error association"),
        ("dialogs", "open invite, question, confirmation and transcript dialogs; test focus trap, initial focus, Escape and return focus"),
        ("status announcements", "trigger upload progress, streaming, processing changes, success and failure alerts"),
        ("color contrast", "measure text, icons, focus, disabled and action buttons including dark-theme controls"),
        ("color independence", "interpret success/warning/failure/selected states without color"),
        ("zoom/reflow", "zoom to 200% and 400% at 1280 CSS pixels"),
        ("text spacing", "apply WCAG text-spacing overrides"),
        ("reduced motion", "enable prefers-reduced-motion and trigger confetti/transitions/loading"),
        ("audio alternatives", "verify transcript availability and non-audio status/content for recordings/responses"),
        ("touch target size", "measure mobile navigation, recorder, player, retry, filters and icon buttons"),
        ("page title and language", "navigate every route and inspect document title/lang"),
    ]
    for title, action in accessibility:
        add("A11Y", f"Accessibility: {title}", "P0" if title in {"keyboard-only navigation", "form labels and errors", "color contrast"} else "P1", "Accessibility",
            "Use representative public, customer and admin data; test with automated scanner plus named assistive technology where applicable.",
            f"1. {action}. 2. repeat in normal, loading, empty, error, disabled, and validation states. 3. record WCAG 2.2 AA criterion and evidence.",
            "All essential information and operations meet WCAG 2.2 AA, remain perceivable/operable/understandable/robust, and any exception has documented impact, owner, workaround, and release decision.", "Automate/Manual", f"WCAG:{title}")
    viewports = ["320×568 mobile", "360×800 Android", "390×844 iPhone", "768×1024 tablet portrait", "1024×768 tablet landscape", "1366×768 laptop", "1440×900 desktop", "1920×1080 desktop"]
    for viewport in viewports:
        add("RWD", f"Responsive layout at {viewport}", "P1", "Responsive UI", "Representative longest labels, validation messages, transcripts, tables, and empty/loading/error states.",
            f"1. render every public/customer/admin route at {viewport}. 2. use navigation, forms, recorder/player, tables, dialogs, filters, and streaming. 3. rotate where relevant.",
            "No unintended horizontal scroll, clipped control, obscured content, overlapping header/footer, unreachable action, or unreadable table occurs; interaction mode remains appropriate to viewport.", "Candidate", f"Viewport:{viewport}")
    browsers = ["Chrome current", "Firefox current", "Safari current", "Edge current", "iOS Safari current", "Android Chrome current"]
    for browser in browsers:
        add("CBR", f"Cross-browser critical journey: {browser}", "P0", "Compatibility", "Supported OS/device with microphone and secure HTTPS origin; Stripe/provider test mode.",
            f"1. on {browser}, complete sign-up/login, checkout test path, dashboard, guided recording/upload, processing status, chat stream/audio, family invitation, settings and logout. 2. inspect console/network/media permissions.",
            "Critical journey and media/SSE/storage/session behavior work without browser-specific blocker; approved minor differences are documented with minimum supported version.", "Manual", f"Browser:{browser}")
    performance = [
        ("public page load", "landing/pricing", "LCP ≤2.5s, INP ≤200ms, CLS ≤0.1 at p75 on agreed mobile profile"),
        ("authenticated dashboard", "profile/recording/processing data", "usable content meets agreed SLO without request waterfall regression"),
        ("admin users list", "10k user dataset with search/filter/pagination", "response/UI remain within agreed p95 and memory bounds"),
        ("question admin", "500 question dataset with reorder/filter", "interaction remains responsive and persistence completes within SLO"),
        ("recording session", "30 answers/long Premium session", "timer/recorder memory remains stable with no progressive leak"),
        ("100 MB upload", "maximum accepted audio payload", "progress remains responsive and server memory/disk stay within capacity"),
        ("chat first token", "concurrent streaming conversations", "p95 time-to-first-token and completion meet product SLO"),
        ("audio streaming", "concurrent range/playback requests", "start/seek remains reliable without buffering regression"),
        ("processing throughput", "queued transcription/voice/analysis workload", "queues drain within capacity target without starvation"),
        ("API soak", "mixed authenticated workload for 4 hours", "latency/error/memory stabilize and no connection/task leak develops"),
    ]
    for title, workload, target in performance:
        add("PERF", f"Performance: {title}", "P0" if title in {"100 MB upload", "processing throughput"} else "P1", "Performance",
            "Production-like sizing, sanitized dataset, warm and cold runs, monitoring for app/DB/Redis/worker/provider timings.",
            f"1. establish baseline. 2. run {workload} at agreed normal, peak, and stress concurrency. 3. measure p50/p95/p99, errors, saturation, and recovery. 4. compare build to baseline.",
            f"{target}; errors remain bounded and explicit; the system recovers after load without manual restart or corrupted/duplicated state.", "Automate", f"NFR:{title}")
    recovery = [
        ("web container restart", "restart Django/Next supervisor container during reads and writes"),
        ("Redis restart", "restart Redis with queued/pending token/task activity"),
        ("Postgres fail/reconnect", "interrupt database pooler during transaction"),
        ("Celery worker restart", "restart each queue worker mid-task"),
        ("disk/storage unavailable", "deny local/Supabase write during upload"),
        ("provider outage", "fail OpenAI/ElevenLabs/Stripe/Brevo for bounded window"),
        ("host reboot", "restart full Docker host after committed and in-flight actions"),
    ]
    for title, action in recovery:
        add("RECOV", f"Recovery: {title}", "P0", "Resilience/DR", "Production-like environment with backups, metrics, correlation IDs, and safe fault-injection approval.",
            f"1. {action}. 2. observe user-facing behavior and alerts. 3. restore dependency. 4. reconcile records/objects/jobs/provider calls. 5. rerun critical smoke.",
            "Committed data remains durable; incomplete work resolves to one clear retryable/terminal state; no duplicate charge/email/provider artifact occurs; service meets agreed RTO/RPO and alerts identify the fault.", "Manual", f"Recovery:{title}")


def add_deployment_operations_cases() -> None:
    deployment = [
        ("Compose default launch", "run docker compose up -d --build from documented production directory", "all required overrides/environment are incorporated by the intended deployment command"),
        ("Compose configuration validation", "run docker compose config and inspect resolved services/networks/labels without printing secrets", "configuration is syntactically valid and uses intended files/values"),
        ("required environment variables", "remove each required secret/URL/provider setting in isolation", "startup fails fast with a named sanitized configuration error"),
        ("optional environment defaults", "omit each documented optional value", "safe documented default is applied consistently"),
        ("Traefik host routing", "request both voicevault-web and web Hostinger hostnames", "Host rules route to the web service without Traefik 404"),
        ("canonical hostname", "open alternate hostname and deep links", "one approved canonical origin is used without cookie/CSRF/CORS split-brain"),
        ("HTTP redirect", "request every hostname/path over port 80", "redirect preserves safe path/query and terminates at HTTPS once"),
        ("certificate lifecycle", "issue, renew, restart and reload Let's Encrypt certificate", "valid certificate persists and renews before expiry"),
        ("Django health check", "call health endpoint before/after dependencies become ready", "health/readiness meaning is accurate for orchestration"),
        ("container dependency startup", "start Postgres/Redis/web/workers in varied order", "services retry/wait safely and converge without manual race workaround"),
        ("database migrations", "deploy forward from supported previous schema with production-like data", "migrations are repeatable, bounded, and preserve data/index integrity"),
        ("static/Next assets", "restart/redeploy while old browser cache exists", "hashed assets and HTML remain compatible without blank page"),
        ("worker queue routing", "enqueue transcription, voice, analysis and default tasks", "each queue has an active consumer and correct routing"),
        ("worker/beat singleton policy", "scale workers and start duplicate beat", "workers scale safely and scheduler duplication is prevented/detected"),
        ("log rotation", "generate sustained web/worker/Traefik logs", "disk usage is bounded and logs remain searchable"),
        ("backup restore", "restore database and audio storage to isolated environment", "documented RPO data is internally consistent and critical smoke passes"),
        ("rollback", "deploy a deliberately failing release then roll back", "previous compatible version restores within RTO without schema/data loss"),
    ]
    for title, action, expected in deployment:
        add("OPS", f"Deployment/operations: {title}", "P0", "Deployment", "Production-like Hostinger host with non-production credentials/data and operator access.",
            f"1. {action}. 2. inspect container health, routes, logs, metrics, persistent volumes, and user-facing response. 3. restart/repeat. 4. run release smoke.",
            f"{expected}; no secret is printed into CI logs or client output; services are reproducible and operator evidence identifies build/config version.", "Automate/Manual", f"OPS:{title}")
    observability = ["web 5xx spike", "database 503 spike", "Celery queue backlog", "processing failure rate", "OpenAI/ElevenLabs 401 or 429", "Stripe webhook failure", "Brevo send failure", "certificate near expiry", "disk/storage near capacity", "security/authentication anomaly"]
    for signal in observability:
        add("OBS", f"Alert and diagnose: {signal}", "P0" if signal in {"Stripe webhook failure", "certificate near expiry", "disk/storage near capacity"} else "P1", "Observability",
            "Monitoring/alerting destination and runbook are configured in staging; synthetic fault can be generated safely.",
            f"1. trigger {signal}. 2. measure detection and notification delay. 3. follow linked dashboard/log correlation/runbook. 4. resolve fault. 5. confirm alert recovery and incident evidence.",
            "Alert fires within agreed threshold with environment/service/severity/correlation context and no secrets; runbook leads to diagnosis; recovery notification occurs without alert storm.", "Candidate", f"Observability:{signal}")


def build_cases() -> None:
    add_route_access_cases()
    add_auth_cases()
    add_consent_profile_cases()
    add_plan_quota_cases()
    add_payment_cases()
    add_dashboard_recording_cases()
    add_processing_cases()
    add_chat_cases()
    add_family_cases()
    add_settings_cases()
    add_admin_cases()
    add_api_contract_cases()
    add_integration_cases()
    add_security_privacy_cases()
    add_nonfunctional_cases()
    add_deployment_operations_cases()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_width(table, total_dxa: int = 9360, indent_dxa: int = 0) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    width = tbl_pr.find(qn("w:tblW"))
    width.set(qn("w:w"), str(total_dxa))
    width.set(qn("w:type"), "dxa")
    ind = tbl_pr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tbl_pr.append(ind)
    ind.set(qn("w:w"), str(indent_dxa))
    ind.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_keep_with_next(paragraph, value=True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "Update this field in Word to refresh."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (fld_char, instr_text, separate, display, end):
        run._r.append(node)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Page ")
    add_field(paragraph, "PAGE")
    paragraph.add_run(" of ")
    add_field(paragraph, "NUMPAGES")


def set_cell_text(cell, text: str, bold=False, color=DARK, size=8.5, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell)


def add_compact_table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]], widths: list[int], font_size=8.5) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    set_table_width(table)
    for i, header in enumerate(headers):
        set_cell_width(table.rows[0].cells[i], widths[i])
        set_cell_shading(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], header, True, DARK_BLUE, 8.5)
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        for i, value in enumerate(values):
            set_cell_width(row.cells[i], widths[i])
            set_cell_text(row.cells[i], str(value), False, DARK, font_size)
            if row_index % 2 == 1:
                set_cell_shading(row.cells[i], PALE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc: Document, items: Iterable[str], level=0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def configure_document(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Title", 34, DARK_BLUE, 0, 12), ("Subtitle", 16, GREY, 0, 10),
        ("Heading 1", 16, BLUE, 18, 10), ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = name != "Subtitle"
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = name.startswith("Heading")

    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.15

    if "Caption Small" not in [s.name for s in styles]:
        cap = styles.add_style("Caption Small", WD_STYLE_TYPE.PARAGRAPH)
        cap.font.name = "Calibri"
        cap.font.size = Pt(8.5)
        cap.font.color.rgb = RGBColor.from_string(GREY)
        cap.paragraph_format.space_after = Pt(4)

    header = sec.header.paragraphs[0]
    header.text = "VOICEVAULT  |  QA MASTER TEST PLAN"
    header.style = styles["Caption Small"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(GREY)

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("VOICEVAULT")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(GOLD)

    p = doc.add_paragraph(style="Title")
    p.add_run("Master QA Test Plan")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("End-to-end customer, administrator, API, AI, payment, security, and operations validation")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("A complete release-quality playbook for the VoiceVault platform")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    add_compact_table(doc, ["Document", "Value"], [
        ("Version", "1.0"),
        ("Prepared", date(2026, 8, 5).strftime("%d %B %Y")),
        ("Product baseline", "VoiceVault web application and Django/Celery backend inspected in the project workspace"),
        ("Coverage", f"{len(CASES)} detailed test cases plus release gates, data, environments, traceability, and reporting"),
        ("Audience", "QA engineers, developers, product owner, security reviewer, operations, and UAT participants"),
        ("Classification", "Internal quality and release-control document; use synthetic test data only"),
    ], [1800, 7560], 10)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Approval status: Draft for baseline review")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(RED)
    doc.add_page_break()


def add_document_control(doc: Document) -> None:
    doc.add_heading("Document control", level=1)
    add_compact_table(doc, ["Role", "Name", "Responsibility", "Sign-off/date"], [
        ("Product owner", "TBD", "Scope, expected behavior, acceptance and release decision", "TBD"),
        ("QA lead", "TBD", "Plan ownership, execution, evidence, defects and quality recommendation", "TBD"),
        ("Engineering lead", "TBD", "Technical readiness, fixes, unit/integration coverage and deployment", "TBD"),
        ("Security/privacy", "TBD", "Security, consent, data handling and retention approval", "TBD"),
        ("Operations", "TBD", "Environment, monitoring, backup, rollback and production readiness", "TBD"),
    ], [1400, 1300, 4850, 1810], 8.8)
    doc.add_heading("Revision history", level=2)
    add_compact_table(doc, ["Version", "Date", "Author", "Change"], [
        ("1.0", "05 Aug 2026", "Codex / project team", "Initial full-project QA baseline derived from current routes, APIs, models, integrations and deployment configuration."),
    ], [900, 1400, 1900, 5160], 8.8)
    doc.add_heading("How to use this plan", level=2)
    add_numbered(doc, [
        "Agree the assumptions, environments, supported platform matrix, retention policy, service-level targets, and acceptance thresholds before execution.",
        "Import or mirror the test cases into the team's test-management system while preserving the case ID and requirement reference.",
        "Execute the P0 smoke set on every deploy, affected P0/P1 regression on every change, and the full regression/security/accessibility packs before production release.",
        "Attach evidence to every execution: build, environment, test-data IDs, request/correlation IDs, screenshots or recordings, logs, and provider test references—never raw secrets or unnecessary personal data.",
        "Record deviations as defects or formally approved exceptions; update this master when routes, entitlements, providers, policy, or deployment topology change.",
    ])
    doc.add_heading("Table of contents", level=2)
    doc.add_paragraph("Major sections are listed below. All numbered headings are also available through Word's Navigation Pane.", style="Caption Small")
    add_compact_table(doc, ["Sections", "Contents"], [
        ("1–5", "Executive summary; system under test; objectives; scope; assumptions and open decisions"),
        ("6–11", "Risks; test levels; environments; test data; entry/exit criteria; severity and defect workflow"),
        ("12.1–12.14", "Public/navigation/authentication; consent/profile/quotas/payments; dashboard; recording; AI processing; chat; family; settings"),
        ("12.15–12.20", "Administrator dashboard, users, questions, processing, payments, logs and audit"),
        ("12.21–12.31", "API and integrations; security/privacy; accessibility/responsive/browser; performance/recovery; deployment/monitoring"),
        ("13–19", "Execution packs; traceability; reporting; automation roadmap; readiness/evidence; final quality decision"),
    ], [1500, 7860], 8.6)


def add_strategy(doc: Document) -> None:
    doc.add_heading("1. Executive summary", level=1)
    doc.add_paragraph(
        f"This plan defines {len(CASES)} detailed, uniquely identified tests for VoiceVault. It validates the complete customer and administrator experience, the Django REST contracts, AI/audio processing, plan enforcement, payments, email and invitations, security/privacy, accessibility, performance, resilience, and Hostinger/Docker/Traefik operations. The emphasis is on release-blocking risks: protecting voice and transcript data, preventing unauthorized family/admin access, avoiding false AI/payment success, enforcing consent and quotas, and recovering safely from asynchronous/provider failures."
    )
    doc.add_heading("2. Product and system under test", level=1)
    doc.add_paragraph(
        "VoiceVault is a web platform in which customers register, purchase or use a plan, answer guided questions by recording audio, upload recordings, and run an asynchronous AI pipeline that transcribes speech, analyzes personality, clones a consented voice, and finalizes an AI identity. Owners and accepted family members can then chat with authorized AIs and, where entitled, hear synthesized responses. Administrators monitor users, questions, processing, payments and logs, and can initiate controlled processing operations."
    )
    add_compact_table(doc, ["Layer", "In-scope implementation"], [
        ("Frontend", "Next.js/React routes for public, authentication, customer, recording, processing, chat, family, settings and admin experiences; client stores and API modules."),
        ("Backend", "Django REST endpoints for users, recordings/questions, AI processing, chat/voice, family invitations/access, payments and admin dashboard."),
        ("Async/data", "Celery worker/beat; Redis broker/cache; PostgreSQL/Supabase database; local or Supabase audio storage."),
        ("External", "OpenAI transcription/personality/chat; ElevenLabs cloning/TTS; Stripe Checkout/webhooks; Brevo SMTP; Traefik and Let's Encrypt."),
        ("Deployment", "Docker Compose services, Gunicorn/Next through supervisor, Traefik host routing/HTTPS, environment-based configuration and health behavior."),
    ], [1700, 7660], 9)
    doc.add_heading("3. Objectives", level=1)
    add_bullets(doc, [
        "Prove that every supported customer and administrator workflow produces the correct visible and persisted outcome.",
        "Prove authentication, role checks, ownership, family grants, consent, quota and plan rules at both UI and API boundaries.",
        "Prove asynchronous processing is ordered, observable, idempotent and recoverable, including the reported personality-analysis failure path.",
        "Prove payment and invitation/email outcomes are verified rather than inferred and remain consistent during retries, replays and provider outages.",
        "Prove voice, transcript, conversation, payment and identity data remain confidential and are handled according to explicit retention/consent rules.",
        "Establish repeatable smoke, regression, UAT and production-readiness gates with evidence and ownership."
    ])
    doc.add_heading("4. Scope", level=1)
    doc.add_heading("4.1 Included", level=2)
    add_bullets(doc, [
        "All discovered public, authentication, customer and administrator routes and their loading, empty, success, error and unauthorized states.",
        "All discovered frontend API clients and backend URL contracts, including object-level authorization and malformed-request behavior.",
        "Free and Premium entitlements: questions 5/30; minutes 15/300; monthly text 5/1000; voice 0/200; family 1/10; storage 500/5000 MB; AI generations 1/3.",
        "Audio recording, local draft, review, validation, multipart upload, storage, transcript/personality/voice/finalization, chat and media playback.",
        "Family invitation creation, Brevo delivery, seven-day single-use token, acceptance, resend, removal, access listing and revocation.",
        "Stripe test-mode checkout, confirmation, webhook verification/idempotency, billing state and administrator reconciliation.",
        "OWASP-oriented security, privacy/consent, WCAG 2.2 AA, responsive/cross-browser, performance, recovery, Docker/Traefik/TLS, monitoring and release controls."
    ])
    doc.add_heading("4.2 Conditional/out of scope until policy is supplied", level=2)
    add_bullets(doc, [
        "Real production charges, real consumer biometric/voice samples, destructive testing against production, and provider load beyond contracted sandbox limits.",
        "Native mobile applications (none were found); mobile web is in scope.",
        "Legal approval of terms/privacy language and jurisdiction-specific compliance. QA verifies implemented behavior against the approved policy once supplied.",
        "Exact account-deletion, transcript/audio retention, refund, password-reset and support/escalation policies where the implementation or requirement is incomplete; these are tracked as release decisions, not silently assumed.",
        "Penetration testing by an independent specialist. This plan includes security functional tests but does not replace an authorized penetration assessment."
    ])
    doc.add_heading("5. Assumptions and open requirement decisions", level=1)
    add_compact_table(doc, ["ID", "Decision required", "QA handling until resolved", "Release impact"], [
        ("A-01", "Canonical public hostname: voicevault-web… or web…", "Test both; require one canonical redirect/origin and aligned CORS/CSRF/cookies.", "P0 if host returns Traefik 404 or origins split."),
        ("A-02", "Password-reset backend and final policy", "Validate non-enumeration/rate limit only where feature is implemented; block release if advertised flow is dead.", "P1/P0 if required for launch."),
        ("A-03", "Account deletion and data retention by data class", "Require approved matrix before destructive/privacy sign-off.", "P0 privacy decision."),
        ("A-04", "Referenced question deletion behavior", "Test block, soft delete or documented cascade; reject silent history corruption.", "P1."),
        ("A-05", "Family member PATCH support", "Treat current client/server mismatch as a contract risk; hide UI or implement endpoint.", "P1 if user-visible."),
        ("A-06", "Performance SLOs and supported browser versions", "Use proposed targets in this plan and baseline current build; product/ops must approve final thresholds.", "P1 readiness."),
        ("A-07", "SMTP sender/domain verification", "Use a Brevo-verified sender. Public Hostinger hostname may be used for application links; sender identity must still satisfy provider policy.", "P0 for invitation/reset deliverability."),
        ("A-08", "Free/Premium commercial wording and one-time/lifetime entitlement", "Reconcile UI, Stripe price, User flags, Payment and package API before release.", "P0 revenue/access."),
    ], [700, 2800, 3950, 1910], 8.4)

    doc.add_heading("6. Risk-based test priorities", level=1)
    add_compact_table(doc, ["Risk", "Failure consequence", "Primary controls/tests"], [
        ("Unauthorized voice/transcript/chat access", "Sensitive personal/biometric-like data disclosure", "P0 role/ownership/IDOR, storage URL, cache and revocation tests"),
        ("False payment success or duplicate grant", "Revenue loss, incorrect entitlement, support burden", "Signed webhook, checkout ownership, replay/order/idempotency, reconciliation"),
        ("AI pipeline false success/duplication", "Unusable AI, provider cost, corrupted state", "Prerequisite/consent gates, task order, retry, concurrency, worker crash"),
        ("Recording loss", "Irreplaceable user effort/audio", "Draft/review, disconnect/refresh, upload retry/idempotency, storage reconciliation"),
        ("Invitation hijack", "Unauthorized family access", "256-bit single-use token, expiry, replay/tamper, resend/revoke and logging"),
        ("Admin privilege misuse", "Cross-account mutation or disclosure", "Allowlist/role enforcement, object checks, confirmations, append-only audit"),
        ("Misconfigured production", "Traefik 404, insecure debug/secrets, failed email/providers", "Compose/host/TLS/env/startup/health/secret scan and smoke"),
    ], [2250, 3000, 4110], 8.7)
    doc.add_heading("7. Test levels and techniques", level=1)
    add_compact_table(doc, ["Level/type", "Purpose", "Typical tooling/evidence"], [
        ("Unit", "Models, serializers, validators, quota arithmetic, token/security helpers, task planning", "pytest/Django TestCase; deterministic provider fakes; branch coverage"),
        ("API/component", "Status/schema/authorization/business rules for every endpoint", "pytest/DRF APIClient; schema assertions; DB/storage/job inspection"),
        ("Integration", "Database, Redis, storage, Celery and sandbox provider contracts", "Docker test environment; provider mocks/sandboxes; request correlation"),
        ("UI/E2E", "Real browser journeys and state/error rendering", "Playwright; accessibility scans; screenshots/video/network trace"),
        ("Exploratory", "Novel sequencing, interruption, usability and audio/device behavior", "Chartered sessions with notes and evidence"),
        ("Security/privacy", "Trust boundaries, OWASP abuse, secrets, consent, retention", "Automated scanners plus manual authorized review; redacted evidence"),
        ("Performance/recovery", "Latency, throughput, saturation, durability and operational recovery", "Load tool, APM/metrics, queue/DB/provider dashboards, fault injection"),
        ("UAT", "Business acceptance and clarity for customers/admins", "Scripted personas; product-owner sign-off"),
    ], [1500, 3800, 4060], 8.7)
    doc.add_heading("8. Environment and configuration matrix", level=1)
    add_compact_table(doc, ["Environment", "Purpose", "Configuration/data", "External services"], [
        ("Local CI", "Fast unit/API/component and frontend checks", "Ephemeral DB/Redis/storage; DEBUG off in security job; synthetic fixtures", "Mocks/fakes; no live billing/email"),
        ("Integration", "Docker/async/storage contract", "Compose topology; migrations; isolated buckets/queues", "Provider mocks or restricted sandboxes"),
        ("Staging/Hostinger", "Release candidate, HTTPS, cross-browser, UAT, provider sandbox", "Production-like DEBUG=False, canonical host, synthetic accounts/audio", "Stripe test; Brevo test recipients; controlled OpenAI/ElevenLabs"),
        ("Production smoke", "Post-deploy non-destructive availability", "Dedicated synthetic monitor account; no destructive/load tests", "Minimal verified transactions only if approved"),
    ], [1350, 2450, 3000, 2560], 8.5)
    doc.add_paragraph("Required configuration review: SECRET_KEY/JWT key; DEBUG; ALLOWED_HOSTS; FRONTEND/NEXT_PUBLIC/BACKEND URLs; CORS/CSRF origins; database/Redis/storage; OpenAI and ElevenLabs models/keys; Stripe mode/price/webhook secret; SMTP host 587/TLS/user/password/from address; Traefik rule/cert resolver; ADMIN_EMAILS; log level and secret redaction. Never place actual credential values in test evidence.")
    doc.add_heading("9. Test data strategy", level=1)
    add_compact_table(doc, ["Persona/fixture", "Minimum data"], [
        ("Anonymous", "No cookies/tokens; public and attack tests."),
        ("New Free customer", "No recordings; default quotas; no consents/AI/payment/family."),
        ("Boundary Free customer", "Usage at limit−1, limit and over-limit for each entitlement."),
        ("Premium customer", "Successful test payment; Premium quotas; valid consent; AI-ready and in-progress variants."),
        ("Family invitee", "Pending, expired, accepted, removed and replayed-token variants."),
        ("Administrator", "Allow-listed; separate second admin for concurrency; protected self/last-admin scenario."),
        ("Audio set", "Valid WebM/MP3/WAV; silence, noise, Unicode speech, short/long, max-size, corrupt, polyglot, MIME mismatch."),
        ("Provider set", "Success, 401, 429, timeout, 5xx, malformed response, delayed response and duplicate callback fixtures."),
        ("Scale set", "10k users; payments/jobs/history; 500 questions; storage/queue load appropriate to performance environment."),
    ], [2200, 7160], 8.7)
    doc.add_paragraph("Use deterministic factories and unique run prefixes. Reset through approved APIs/factories, not shared manual data. Synthetic audio must not imitate a real person without consent. Preserve IDs required for defect reproduction, then purge test data according to environment policy.")
    doc.add_heading("10. Entry, suspension, resumption and exit criteria", level=1)
    doc.add_heading("10.1 Entry", level=2)
    add_bullets(doc, [
        "Release candidate is versioned and deployed; migrations complete; health checks pass; required queues and providers are configured in test mode.",
        "Requirements/open decisions are approved or explicitly waived; test accounts, audio and provider fixtures are available; monitoring/evidence access works.",
        "No unresolved environment blocker prevents meaningful execution; previous release's critical defects have verified disposition."
    ])
    doc.add_heading("10.2 Suspend and resume", level=2)
    doc.add_paragraph("Suspend affected execution when the environment is materially unstable, test data is corrupted, credentials/providers are unintentionally production-scoped, or a defect prevents more than 25% of planned cases in a critical area. Record scope, evidence, owner and time. Resume after a verified fix/configuration correction, environment smoke, and data re-baseline; do not mark blocked cases failed solely because of an environment incident.")
    doc.add_heading("10.3 Exit/release gate", level=2)
    add_bullets(doc, [
        "100% of P0 and ≥95% of planned P1 cases executed on the release candidate; all critical customer/admin journeys pass.",
        "No open Severity 1 or Severity 2 defect; any lower defect has owner, target, accepted customer/security impact and product approval.",
        "No unresolved unauthorized access, false payment/AI success, recording-loss, secret-exposure, consent-bypass, Traefik/TLS, migration/rollback or critical accessibility defect.",
        "Performance/capacity, backup restore, monitoring alerts, security review, accessibility assessment and UAT meet approved thresholds.",
        "Test summary, defect list, residual risks, rollback decision and named sign-offs are stored with build/environment evidence."
    ])
    doc.add_heading("11. Severity, priority and defect workflow", level=1)
    add_compact_table(doc, ["Severity", "Definition", "Examples", "Release rule"], [
        ("S1 Critical", "Security/data loss, incorrect charge/access, total critical outage, no workaround", "Cross-user audio; secret leak; paid state without verified payment; irrecoverable recordings", "Immediate stop; must fix and fully regress"),
        ("S2 High", "Major function broken for many users or reliable workaround absent", "Pipeline cannot complete; login/record/upload blocked; admin cannot operate", "Must fix before release"),
        ("S3 Medium", "Partial/noncritical impact with workaround", "One filter/status/error path wrong; layout impairs subset", "Fix or written acceptance"),
        ("S4 Low", "Cosmetic/minor copy or low-risk inconsistency", "Spacing/copy/rare visual issue", "May defer with owner"),
    ], [1100, 3000, 3300, 1960], 8.5)
    doc.add_paragraph("Priority indicates how urgently the test must run: P0 on every relevant deploy and before release; P1 in affected/full regression; P2 planned usability/compatibility/edge coverage. Defect lifecycle: New → Triaged → In progress → Ready for QA → Retest → Closed, with Reopened/Deferred/Duplicate/Not a defect only after evidence. Every defect records build, environment, role, data IDs, exact steps, actual/expected, reproducibility, impact, evidence, logs/correlation IDs and regression scope.")
    doc.add_page_break()


AREA_NAMES = {
    "PUB": "Public pages", "NAV": "Navigation and route protection", "ADM-AUTH": "Administrator authentication and authorization",
    "AUTH": "Authentication and session management", "CNS": "Consent", "PRO": "Profile", "QTA": "Plan entitlements and quotas",
    "PAY": "Payments and pricing", "DSH": "Customer dashboard", "REC": "Recording and uploads", "AIP": "AI processing",
    "CHT": "Chat, voice input and AI audio", "FAM": "Family invitations and access", "SET": "Settings",
    "ADM": "Admin dashboard", "USR": "Admin user management", "QUE": "Admin question management",
    "AADM": "Admin processing operations", "APAY": "Admin payment operations", "LOG": "Admin logs and audit",
    "API": "API contracts and platform behavior", "INT": "Third-party and infrastructure integrations",
    "SEC": "Application security", "PRV": "Privacy and data governance", "A11Y": "Accessibility",
    "RWD": "Responsive UI", "CBR": "Cross-browser compatibility", "PERF": "Performance and capacity",
    "RECOV": "Reliability and recovery", "OPS": "Deployment and operations", "OBS": "Monitoring and alerting",
}


def add_test_catalog(doc: Document) -> dict[str, list[tuple[str, TestCase]]]:
    doc.add_heading("12. Detailed test-case catalog", level=1)
    doc.add_paragraph("Execution notation: Preconditions establish the required fixture. Procedures are numbered in the order performed and include UI/API/persistence verification. Expected results are pass criteria, not suggestions. 'Automate' indicates stable regression value; 'Candidate' requires implementation assessment; 'Manual' requires human/device/operational judgment; combined labels require both.")
    grouped: dict[str, list[tuple[str, TestCase]]] = {}
    counts: dict[str, int] = {}
    for case in CASES:
        counts[case.area] = counts.get(case.area, 0) + 1
        case_id = f"{case.area}-{counts[case.area]:03d}"
        grouped.setdefault(case.area, []).append((case_id, case))
    area_order = list(AREA_NAMES)
    section_no = 1
    for area in area_order:
        items = grouped.get(area, [])
        if not items:
            continue
        doc.add_heading(f"12.{section_no} {AREA_NAMES[area]} ({len(items)} cases)", level=2)
        rows = []
        for case_id, c in items:
            descriptor = c.title + "\nPreconditions: " + c.preconditions
            procedure = c.steps
            expected = c.expected + (f"\nEvidence/automation: {c.automation}." if c.automation else "")
            rows.append((case_id, c.priority, c.test_type, descriptor, procedure, expected))
        add_compact_table(doc,
                          ["ID", "Pri", "Type", "Test and preconditions", "Procedure", "Expected result"],
                          rows,
                          [700, 450, 1050, 2100, 2700, 2360],
                          7.7)
        section_no += 1
    return grouped


def add_execution_packs(doc: Document, grouped: dict[str, list[tuple[str, TestCase]]]) -> None:
    doc.add_page_break()
    doc.add_heading("13. Execution packs", level=1)
    p0s = [(case_id, c) for area in grouped.values() for case_id, c in area if c.priority == "P0"]
    doc.add_heading("13.1 Per-deployment smoke", level=2)
    doc.add_paragraph("The automated smoke should finish quickly and use non-destructive synthetic data. At minimum it verifies:")
    add_bullets(doc, [
        "HTTPS canonical host and health; landing/login; administrator route denial for customer.",
        "Valid login/profile; dashboard load; active question list; small valid audio upload; processing status retrieval.",
        "Provider-mocked transcription/personality/voice/finalization; AI-ready chat streaming; family access authorization.",
        "Package retrieval; Stripe-signed webhook idempotency; Brevo SMTP connectivity in staging; logout and old-session denial.",
    ])
    doc.add_paragraph(f"Full P0 regression contains {len(p0s)} cases. Generate the run by filtering Priority=P0 in the catalog; no P0 failure may be waived without named product, engineering and security/operations approval appropriate to the risk.")
    doc.add_heading("13.2 Change-based regression", level=2)
    add_compact_table(doc, ["Changed area", "Mandatory regression"], [
        ("Authentication/user/profile", "AUTH, NAV, ADM-AUTH, CNS, PRO, authorization subset of API/SEC, protected route smoke"),
        ("Plan/price/payment", "QTA, PAY, APAY, Stripe INT/API/SEC, dashboard/settings and full entitlement checks"),
        ("Recording/questions/storage", "REC, QUE, storage INT, quota, processing prerequisites, mobile/browser audio"),
        ("AI tasks/providers", "AIP, AADM, OpenAI/ElevenLabs INT, CHT, queue recovery/performance/observability"),
        ("Chat/family", "CHT, FAM, invitation email, object authorization, quotas, privacy/revocation"),
        ("Frontend styling/components", "Affected functional cases plus A11Y/RWD/CBR and visual states across customer/admin"),
        ("Infrastructure/config", "OPS, OBS, RECOV, security headers/secrets, critical end-to-end smoke and rollback"),
    ], [2500, 6860], 8.7)
    doc.add_heading("13.3 UAT journeys", level=2)
    add_numbered(doc, [
        "New customer: understand offering → sign up → see Free limits → record guided answers → review/upload → grant explicit AI consents → observe progress → chat with ready AI.",
        "Premium buyer: pricing clarity → Stripe test checkout → verified success → Premium limits visible → extended recording/voice response/family capacity works.",
        "Family owner and invitee: send branded invite → accept single-use link → invitee signs in and chats only with authorized AI → owner removes access → access immediately ends.",
        "Administrator: authenticate → reconcile statistics → find a user → manage questions/order → diagnose a failed personality step → retry safely → reconcile payment and audit trail.",
        "Operations: deploy with one command/configured environment → verify canonical HTTPS/no Traefik 404 → verify worker queues/SMTP/providers → trigger alert → restore/rollback and pass smoke."
    ])
    doc.add_heading("14. Requirements traceability summary", level=1)
    rows = []
    for area in AREA_NAMES:
        items = grouped.get(area, [])
        if items:
            ids = f"{items[0][0]}–{items[-1][0]}"
            rows.append((AREA_NAMES[area], ids, str(len(items)), "P0/P1/P2 as listed"))
    add_compact_table(doc, ["Requirement/capability", "Case range", "Count", "Release use"], rows, [3600, 2000, 900, 2860], 8.3)
    doc.add_heading("15. Execution reporting and metrics", level=1)
    add_bullets(doc, [
        "Daily: planned/executed/pass/fail/blocked/not-run by priority and area; new/open/closed defects by severity; environment incidents; top risks and next decisions.",
        "Release: P0/P1 completion and pass rate; defect leakage/reopen rate; requirements coverage; automation pass/stability; accessibility/security/performance exceptions; UAT/sign-off status.",
        "Operational quality: API p95/error rate, provider error/rate-limit/cost, queue age/failures/retries, upload success, processing completion time, chat first token, email acceptance, webhook failures, certificate/disk capacity.",
        "Do not use pass percentage alone: publish untested scope, blocked cases, waived defects and residual risk in plain language."
    ])
    doc.add_heading("16. Automation roadmap", level=1)
    add_compact_table(doc, ["Phase", "Deliverables", "Quality bar"], [
        ("1 — release protection", "P0 API authorization/validation; auth/session; quota; Stripe webhook; AI prerequisite/idempotency; Playwright smoke", "Runs in CI; deterministic; no real provider spend; failing evidence retained"),
        ("2 — core regression", "Recording/upload matrices; customer/admin journeys; family tokens; chat stream parser; question CRUD/reorder", "Parallel-safe fixtures; provider contract fakes; <2% flaky rerun rate"),
        ("3 — nonfunctional", "axe accessibility; visual/responsive; load baselines; dependency recovery; secret/config/container scans", "Thresholds versioned; production-like scheduled pipeline; trend reporting"),
        ("4 — production assurance", "Synthetic HTTPS/login/read-only chat health, queue/provider/cert/storage alerts", "No personal data; bounded cost; on-call runbooks and alert ownership"),
    ], [1600, 5000, 2760], 8.6)
    doc.add_paragraph("Current automated coverage discovered in the repository is limited: one frontend customer-flow E2E specification and a small AI-processing backend test module. That is a useful seed but materially below this release plan; automation should begin with authorization, idempotency, quotas, AI prerequisites, and critical smoke rather than attempting to automate every visual/device case first.")
    doc.add_heading("17. Release readiness checklist", level=1)
    checklist = [
        "[ ] Build/commit, environment, migration and configuration baseline recorded",
        "[ ] Canonical Hostinger hostname, DNS, Traefik router and trusted TLS certificate verified",
        "[ ] DEBUG=False; required secrets rotated and server-only; repository/image/client/log scan clean",
        "[ ] Database/Redis/storage/web/worker/beat health and queue routing verified",
        "[ ] Brevo verified sender/TLS authentication/delivery and invitation link confirmed",
        "[ ] OpenAI and ElevenLabs configured model/consent/quota/error handling verified",
        "[ ] Stripe test price, checkout, signed webhook, replay and reconciliation verified",
        "[ ] Migrations, backup restore and rollback rehearsal passed",
        "[ ] P0/P1 execution thresholds met; no open S1/S2; residual risks approved",
        "[ ] Security/privacy, WCAG, browser/mobile, performance and recovery evidence approved",
        "[ ] Customer and administrator UAT signed; support/admin runbooks and alert owners ready",
        "[ ] Production smoke and post-release monitoring window scheduled with rollback authority",
    ]
    add_bullets(doc, checklist)
    doc.add_heading("18. Evidence templates", level=1)
    doc.add_heading("18.1 Test execution record", level=2)
    add_compact_table(doc, ["Field", "Required value"], [
        ("Case/build/environment", "Case ID; version/commit/image; environment and time zone"),
        ("Actor/data", "Role and synthetic user/recording/payment/job IDs; never raw password/token/audio unless secured evidence is essential"),
        ("Result", "Pass / Fail / Blocked / Not run with executor and timestamp"),
        ("Evidence", "Screenshot/video, request/response schema excerpt, correlation/task/provider test ID, relevant redacted log/metric"),
        ("Deviation", "Defect ID or approved exception; actual vs expected; reproducibility and affected scope"),
    ], [2100, 7260], 8.7)
    doc.add_heading("18.2 Defect report minimum", level=2)
    add_bullets(doc, [
        "Concise title containing area and failed outcome; severity/priority with customer/security/business justification.",
        "Exact preconditions, data IDs, ordered reproduction steps, expected and actual results, reproducibility and first known build.",
        "Browser/device or API request context; screenshots/video; redacted logs, trace/correlation/task/provider test IDs.",
        "Affected roles/plans/routes/providers; workaround; data/privacy/payment risk; proposed regression scope and owner."
    ])
    doc.add_heading("19. Final quality decision", level=1)
    doc.add_paragraph("The QA lead issues one of: GO (all gates met); CONDITIONAL GO (no S1/S2, explicit low/medium residual risks with owners and dates); or NO-GO (a gate is unmet or evidence is insufficient). The decision must name the exact release build and environment and be co-signed by product and engineering, with security/privacy/operations participation for risks in their domains.")


def main() -> None:
    build_cases()
    doc = Document()
    configure_document(doc)
    core = doc.core_properties
    core.title = "VoiceVault Master QA Test Plan"
    core.subject = "End-to-end customer and administrator QA plan"
    core.author = "VoiceVault Project Team"
    core.keywords = "VoiceVault, QA, test plan, test cases, admin, customer, API, security, AI"
    add_cover(doc)
    add_document_control(doc)
    add_strategy(doc)
    grouped = add_test_catalog(doc)
    add_execution_packs(doc, grouped)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Generated {OUTPUT}")
    print(f"Test cases: {len(CASES)}")


if __name__ == "__main__":
    main()
