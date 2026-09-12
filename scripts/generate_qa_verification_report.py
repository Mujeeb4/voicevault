from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
QA_DIR = ROOT / "artifacts" / "qa"
LIVE_API = QA_DIR / "live_admin_api_results_2026-08-23.json"
LIVE_UI = QA_DIR / "live_admin_ui_results_2026-08-23.json"
OUTPUT = QA_DIR / "VoiceVault_Independent_QA_Verification_Report_2026-08-23.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17212B"
MUTED = "5B6570"
LIGHT_BLUE = "E8EEF5"
LIGHT_GREY = "F2F4F7"
PALE = "F7F9FB"
WHITE = "FFFFFF"
GREEN = "2E7D32"
GREEN_FILL = "EAF5EC"
RED = "9B1C1C"
RED_FILL = "FCEBEC"
AMBER = "7A5A00"
AMBER_FILL = "FFF5D6"
PURPLE = "5B3E96"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run(run, *, size=11, bold=False, color=INK, italic=False, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)
    return run


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa: list[int]):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_cell_text(cell, text, *, bold=False, color=INK, size=9.2, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    set_run(p.add_run(str(text) if text is not None else ""), size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths_dxa, status_col=None, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        shade(table.rows[0].cells[i], DARK_BLUE)
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=WHITE, size=9.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_data):
            if len(table.rows) % 2 == 1:
                shade(row.cells[i], PALE)
            if status_col is not None and i == status_col:
                value_text = str(value)
                if value_text == "Pass":
                    shade(row.cells[i], GREEN_FILL)
                    color = GREEN
                elif value_text in {"Fail", "No-Go", "High", "Critical"}:
                    shade(row.cells[i], RED_FILL)
                    color = RED
                else:
                    shade(row.cells[i], AMBER_FILL)
                    color = AMBER
                set_cell_text(row.cells[i], value_text, bold=True, color=color, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                set_cell_text(row.cells[i], value, size=font_size, align=align)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for node in (fld_begin, instr, fld_sep, text, fld_end):
        run._r.append(node)


def configure_doc(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Title", 28, INK, 0, 6),
        ("Subtitle", 13, MUTED, 0, 14),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    set_run(hp.add_run("VOICEVAULT  |  INDEPENDENT QA VERIFICATION"), size=8.5, bold=True, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(fp.add_run("Confidential QA evidence  |  Page "), size=8.5, color=MUTED)
    add_page_number(fp)


def add_label_paragraph(doc, label, value, *, color=INK, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    set_run(p.add_run(label + ": "), bold=True, color=DARK_BLUE)
    set_run(p.add_run(value), color=color)
    return p


def add_callout(doc, label, text, kind="info"):
    fill, accent = {
        "info": (LIGHT_BLUE, DARK_BLUE),
        "risk": (RED_FILL, RED),
        "warn": (AMBER_FILL, AMBER),
        "pass": (GREEN_FILL, GREEN),
    }[kind]
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(label + "  "), size=10.5, bold=True, color=accent)
    set_run(p.add_run(text), size=10.5, color=INK)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_major(doc, title):
    p = doc.add_paragraph(title, style="Heading 1")
    p.paragraph_format.page_break_before = True
    return p


def add_case_detail(doc, item):
    p = doc.add_paragraph(style="Heading 3")
    set_run(p.add_run(f"{item['case_id']}  |  {item['title']}"), size=12, bold=True, color=DARK_BLUE)
    add_label_paragraph(doc, "Area", item["area"], after=2)
    add_label_paragraph(doc, "Expected", item["expected"], after=2)
    concise_evidence = {
        "ADM-LIVE-026": "The endpoint returned internal_error for page=abc instead of rejecting the invalid page value.",
        "ADM-LIVE-069": "The endpoint returned 'Questions reordered successfully' with updated_count=0 for a nonexistent UUID.",
        "ADM-LIVE-075": "The audit endpoint returned user and batch-processing events, but no entry referenced the created, updated, reordered, activated, or deleted QA questions.",
    }.get(item["case_id"], item["evidence"])
    add_label_paragraph(doc, "Observed", f"HTTP {item['actual_status']} - {concise_evidence}", color=RED if item["result"] == "Fail" else INK, after=6)


def build_report():
    api = json.loads(LIVE_API.read_text(encoding="utf-8"))
    ui = json.loads(LIVE_UI.read_text(encoding="utf-8")) if LIVE_UI.exists() else None
    doc = Document()
    configure_doc(doc)

    # Memo masthead title block.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run("INDEPENDENT QA VERIFICATION REPORT"), size=10, bold=True, color=BLUE)
    p = doc.add_paragraph(style="Title")
    set_run(p.add_run("VoiceVault"), size=30, bold=True, color=INK)
    p = doc.add_paragraph(style="Subtitle")
    set_run(p.add_run("Workbook reliability audit, live administrator verification, and release-risk assessment"), size=13, color=MUTED)
    add_table(doc, ["Report date", "Environment", "Revision", "Decision"], [["23 Aug 2026", "Hostinger production-like deployment", "f59844354a82", "CONDITIONAL NO-GO"]], [1800, 3000, 1900, 2660], status_col=3, font_size=9.4)

    doc.add_heading("Executive verdict", level=1)
    add_callout(doc, "Release recommendation", "Do not use the supplied workbook as release sign-off. The product is reachable and most live administrator operations work, but the evidence set contains contradictory status labels, unsupported automated claims, missing audit controls, a server-error input path, and unresolved customer-side failures.", "risk")
    summary_rows = [
        ["Friend's workbook", "538", "504 Pass / 20 Fail / 6 Blocked / 8 Skip", "Low confidence"],
        ["Live admin API", "49", "43 Pass / 3 Fail / 3 deferred", "High confidence"],
        ["Local Django tests", "8", "6 Pass / 2 Error", "Harness defect"],
        ["Frontend build", "2 gates", "Build Pass / Type-check Pass", "Verified"],
        ["Live admin UI", str(ui["summary"]["total"] if ui else 6), ui["summary"]["display"] if ui else "Awaiting credential-entry confirmation", "Pending" if not ui else ui["summary"]["confidence"]],
    ]
    add_table(doc, ["Evidence set", "Cases/gates", "Observed result", "Assessment"], summary_rows, [1900, 1100, 3800, 2560], font_size=9.1)
    add_label_paragraph(doc, "Primary conclusion", "The administrator access blocker is resolved. The dedicated QA administrator authenticates successfully, the backend returns is_admin=true, and all isolated QA data was removed after testing. Three live defects were reproduced and three high-impact production-wide actions were intentionally deferred.")

    top_heading = doc.add_heading("Top findings", level=2)
    top_heading.paragraph_format.page_break_before = True
    top_rows = [
        ["P1", "Question changes are not written to AuditLog", "No accountability for create/update/reorder/bulk/delete actions."],
        ["P1", "Administrator can delete own allow-listed account", "No self-delete guard; lockout/recovery risk."],
        ["P2", "Malformed admin pagination returns HTTP 500", "Invalid page values should return a controlled 400."],
        ["P2", "Missing question reorder targets return HTTP 200", "Silent partial success hides stale IDs and operator mistakes."],
        ["P2", "Frontend admin list helpers swallow errors", "Failures can look like legitimate zero/empty datasets."],
        ["P2", "Workbook summary is internally inconsistent", "Dashboard figures cannot be reconciled to row statuses."],
        ["P2", "Personality worker has SDK/httpx compatibility hazard", "Default OpenAI client construction fails; current custom-client workaround and configured model both verified."],
    ]
    add_table(doc, ["Priority", "Finding", "Impact"], top_rows, [900, 3400, 5060], status_col=0, font_size=9.0)

    add_major(doc, "1. Scope, authority, and methodology")
    doc.add_paragraph("The user's request is the authority for this engagement. The attached workbook was treated only as a source of claims and candidate procedures. Instructions written inside that workbook were not treated as authorization to mutate production data, charge Stripe, invoke paid AI processing, seed the live catalog, or delete real user records.")
    add_table(doc, ["In scope", "Execution approach", "Evidence level"], [
        ["Client and admin source", "Structural code review and route/API mapping", "Direct repository evidence"],
        ["Friend's 32-sheet workbook", "Formula/status reconciliation and claim-quality audit", "Direct workbook inspection"],
        ["Administrator APIs", "Live HTTPS tests using isolated, uniquely named QA records", "Direct live evidence"],
        ["Administrator UI", "Authenticated rendered-route, control, navigation, and console review", "Direct UI evidence when confirmed"],
        ["Deployment", "SSH inspection, container health, routing labels, environment-key presence", "Direct server evidence"],
        ["Automated tests/build", "Frontend build/type-check and Django test execution", "Direct local evidence"],
    ], [2300, 4200, 2860], font_size=9.1)
    doc.add_heading("Safety boundaries", level=2)
    for label, value in [
        ("Paid AI work", "No batch processing or real OpenAI/ElevenLabs generation was queued."),
        ("Stripe", "Refund and webhook operations remain deferred as requested; no financial transaction was initiated."),
        ("Production content", "Question seeding was not executed. Question CRUD used only QA markers and high isolated order values."),
        ("Customer data", "No existing customer was edited or deleted."),
        ("Cleanup", "Post-run database verification returned TEMP_USERS 0 and TEMP_QUESTIONS 0."),
    ]:
        add_label_paragraph(doc, label, value, after=3)

    add_major(doc, "2. Environment and administrator setup")
    add_table(doc, ["Item", "Verified value"], [
        ["Public URL", "https://voicevault-web.srv1183929.hstgr.cloud"],
        ["Server", "Hostinger VPS, Ubuntu 24.04.3 LTS, 72.61.76.103"],
        ["Deployment", "Docker Compose: web, celery-worker, celery-beat, PostgreSQL, Redis"],
        ["Code revision", "f59844354a8258bd6371210e9e49e172e4840a41 (local and remote master)"],
        ["Live health", "voicevault-web healthy; /api/users/health/ reports healthy and database connected"],
        ["Administrator model", "Email allowlist through ADMIN_EMAILS; not Django is_superuser/is_staff"],
        ["QA administrator", "qa.superadmin@voicevault.test; backend admin identity verified"],
        ["Runtime change", "ADMIN_EMAILS persisted in .env; web container alone recreated using all Traefik overlays"],
        ["Routing preservation", "Both Host rules and letsencrypt resolver verified after recreation"],
    ], [2300, 7060], font_size=9.3)
    add_callout(doc, "Security note", "The password is intentionally excluded from this report and should be delivered separately. Because authorization is email-based, removing the email from ADMIN_EMAILS and recreating the web service revokes admin access.", "warn")

    add_major(doc, "3. Audit of the supplied QA workbook")
    doc.add_heading("3.1 Status reconciliation", level=2)
    add_table(doc, ["Metric", "Row-level calculation", "Summary-sheet display", "Assessment"], [
        ["Total cases", "538", "538", "Matches"],
        ["Pass", "504", "456", "Mismatch of 48"],
        ["Fail", "20", "20", "Matches"],
        ["Blocked", "6", "Not represented", "Missing category"],
        ["Skip", "8", "Not represented", "Missing category"],
        ["Not tested", "At least 14 non-pass/non-fail", "0", "False completeness"],
    ], [1700, 2100, 2300, 3260], font_size=9.1)
    add_callout(doc, "Material formula defect", "The 48-case Chat sheet is marked Pass at row level but contributes zero passed cases to the Summary. Blocked and Skip are also omitted from the summary logic, so the headline completion view is not a reliable roll-up.", "risk")
    doc.add_heading("3.2 Execution provenance", level=2)
    add_table(doc, ["Tested by", "Rows", "Share", "Confidence note"], [
        ["Automated (Antigravity)", "478", "88.8%", "No reproducible logs, scripts, screenshots, timestamps, or assertions attached"],
        ["Manual (Hanzila)", "23", "4.3%", "Some useful observations; evidence still mostly narrative"],
        ["Blank", "37", "6.9%", "No tester attribution"],
    ], [2500, 900, 900, 5060], font_size=9.2)
    doc.add_paragraph("Automated labels were not accepted as proof. A passing label was considered trustworthy only when the Actual Result described the same behavior as the Expected Result and the procedure was realistically automatable with the evidence supplied.")

    doc.add_heading("3.3 Non-pass rows requiring follow-up", level=2)
    nonpass_rows = [
        ["NAV-002..006", "Fail", "Anonymous protection works, but post-login deep links return to /dashboard instead of /record, /processing, /chat, /family, or /settings."],
        ["ADM-AUTH-001..006", "Blocked", "No administrator was available. This blocker is now resolved and re-tested live."],
        ["AUTH-001..011", "Fail", "Signup/login scenarios were grouped as failures; AUTH-001 cites a payment-service error during signup, which is a flow coupling concern."],
        ["AUTH-012", "Fail", "No Actual Result, tester, date, or defect note; failure is unsupported."],
        ["AUTH-022..028", "Skip", "Refresh-token/API-security cases were not executed."],
        ["AUTH-031", "Skip", "JWT expiry-boundary case was not executed."],
        ["PRO-003", "Fail", "Email could be altered through profile PATCH without verification."],
        ["API-021", "Fail", "Unauthenticated audio request returned 401 while the test expected 400/404; the expectation may be wrong because auth should fail first."],
        ["API-023", "Fail", "Checkout creation returned HTTP 500 instead of 200/400."],
    ]
    add_table(doc, ["Case IDs", "Workbook status", "Audit observation"], nonpass_rows, [1700, 1400, 6260], status_col=1, font_size=8.9)

    doc.add_heading("3.4 Contradictory or unsupported pass claims", level=2)
    contradiction_rows = [
        ["Admin functional sheets 12.15-12.20", "69 Pass rows", "First rows often prove only that a non-admin gets 403; remaining rows repeat generic 'rule verified' text. No admin existed."],
        ["PAY-002", "Pass", "Actual Result says valid checkout returned HTTP 500 Payment service error."],
        ["PAY-003..007", "Pass", "Invalid-package scenarios returned HTTP 500 but were still marked Pass."],
        ["PERF-003", "Pass", "Expected admin users/10k dataset; actual result discusses Questions catalog latency."],
        ["PERF-004", "Pass", "Expected question-admin performance; actual result discusses profile/quota latency."],
        ["Accessibility", "Automated Pass", "Screen-reader/manual usability claims have no assistive-technology evidence."],
        ["Cross-browser/device", "Automated Pass", "Chrome, Firefox, Safari, Edge, iOS, and Android claims have no browser/device matrix artifacts."],
        ["DR/monitoring/deployment", "Automated Pass", "Destructive and operational scenarios use generic text with no runbook logs or recovery evidence."],
    ]
    add_table(doc, ["Area", "Claim", "Why confidence is low"], contradiction_rows, [2400, 1300, 5660], font_size=8.8)

    add_major(doc, "4. Live administrator API verification")
    s = api["summary"]
    add_callout(doc, "Live result", f"{s['executed']} cases executed: {s['passed']} passed and {s['failed']} failed. {s['not_executed']} additional cases were explicitly deferred or unavailable. All mutable tests used isolated QA records and cleanup was verified.", "info")
    rows = []
    for item in api["results"]:
        observed = f"HTTP {item['actual_status']}" if item["actual_status"] is not None else item["execution"]
        rows.append([item["case_id"], item["area"], item["title"], observed, item["result"]])
    add_table(doc, ["ID", "Area", "Test", "Observed", "Result"], rows, [1500, 1700, 3560, 1200, 1400], status_col=4, font_size=8.2)

    doc.add_heading("4.1 Reproduced failures", level=2)
    for item in [x for x in api["results"] if x["result"] == "Fail"]:
        add_case_detail(doc, item)
    doc.add_heading("4.2 Deferred cases", level=2)
    deferred_rows = [[x["case_id"], x["title"], x["execution"], x["evidence"]] for x in api["results"] if x["result"] == "Not Executed"]
    add_table(doc, ["ID", "Case", "Disposition", "Reason"], deferred_rows, [1500, 2500, 1500, 3860], font_size=8.8)

    add_major(doc, "5. Administrator UI verification")
    if ui:
        add_callout(doc, "UI result", ui["summary"]["display"], "pass" if ui["summary"].get("failed", 0) == 0 else "risk")
        ui_rows = [[x["route"], x["title"], x["result"], x["observation"]] for x in ui["results"]]
        add_table(doc, ["Route", "Screen", "Result", "Observation"], ui_rows, [1700, 1900, 1200, 4560], status_col=2, font_size=8.8)
        if ui.get("console_errors"):
            doc.add_heading("Console/runtime evidence", level=2)
            for error in ui["console_errors"]:
                add_label_paragraph(doc, error.get("route", "Route"), error.get("message", ""), color=RED)
    else:
        add_callout(doc, "Pending", "The authenticated browser pass awaits action-time confirmation to enter the generated QA credential into the VoiceVault login form. API authorization and all admin endpoints were verified independently.", "warn")

    add_major(doc, "6. Personality-analysis failure investigation")
    add_table(doc, ["Evidence", "Observation", "Interpretation"], [
        ["ProcessingQueue", "Five personality records: four failed at two-minute intervals and one later completed", "Retry/history exists, but stored failure text is generic"],
        ["Stored error", "Personality analysis failed", "Insufficient for root-cause diagnosis"],
        ["Default client construction", "TypeError: Client.__init__() got an unexpected keyword argument 'proxies'", "OpenAI SDK/httpx dependency incompatibility reproduced in worker"],
        ["Task workaround", "Custom httpx.Client passed into openai.OpenAI", "Current code avoids the reproduced constructor failure"],
        ["Configured model", "gpt-4.1-mini", "Read-only model retrieval returned HTTP 200; API key and model are currently valid"],
    ], [2100, 3600, 3660], font_size=9.0)
    add_callout(doc, "Root-cause assessment", "The historical failures cannot be proven from the sanitized July records because the original exception class was not persisted. However, a concrete SDK/httpx constructor incompatibility is reproducible in the live worker, while the current task's custom-client workaround and gpt-4.1-mini availability both pass. Treat the dependency mismatch as the leading historical cause, not a mathematically certain attribution.", "warn")
    add_label_paragraph(doc, "Recommended fix", "Pin a compatible openai/httpx pair, add a startup smoke check that constructs the same client used by the task, persist sanitized exception class/provider request ID, and add an integration test for valid JSON response parsing.")

    add_major(doc, "7. Code-review and test-automation findings")
    code_rows = [
        ["VV-SEC-001", "High", "Admin self-deletion", "admin_user_detail DELETE has no guard against deleting request.admin_user."],
        ["VV-AUD-001", "High", "Question audit gap", "Create/update/delete/reorder/bulk question operations do not create AuditLog records."],
        ["VV-VAL-001", "Medium", "Weak admin user validation", "PATCH blindly applies allow-listed fields; no explicit enum/type validation or concurrency/version check."],
        ["VV-DAT-001", "Medium", "Non-atomic reorder", "Questions are updated one at a time; missing IDs are silently skipped and duplicate/invalid order inputs are accepted."],
        ["VV-DAT-002", "Medium", "Bulk-delete safeguards", "No confirmation contract or historical-reference safety check is enforced server-side."],
        ["VV-PAY-001", "Medium", "Refund unsupported", "Admin payments endpoint is read-only; workbook refund claims do not match implementation."],
        ["VV-FE-001", "Medium", "Error masking", "Several frontend admin API helpers catch errors and return empty arrays or zero KPIs, making outages resemble valid empty states."],
        ["VV-NAV-001", "Medium", "Admin deep-link loss", "AdminRoute redirects every anonymous admin subroute to /login?redirect=/admin rather than preserving the requested route."],
        ["VV-TST-001", "Medium", "Django test connection closure", "8 discovered tests produced 6 passes and 2 errors; the Celery eager-task path closes the TestCase database connection."],
        ["VV-CFG-001", "Low", "SQLite incompatibility", "PostgreSQL-only connection OPTIONS are applied to SQLite DATABASE_URL, preventing lightweight test execution."],
    ]
    add_table(doc, ["Defect", "Severity", "Title", "Evidence"], code_rows, [1300, 1100, 2300, 4660], status_col=1, font_size=8.7)
    doc.add_heading("Verified engineering gates", level=2)
    add_table(doc, ["Gate", "Result", "Evidence"], [
        ["Frontend production build", "Pass", "All 28 pages, including six admin routes, built successfully."],
        ["Frontend TypeScript check", "Pass", "Passed when rerun after build completed."],
        ["Django tests", "Error", "6/8 passed; 2 errors tied to closed database connection in eager Celery task test."],
        ["Live health", "Pass", "Web container healthy; database and public health endpoint responsive."],
    ], [2600, 1300, 5460], status_col=1, font_size=9.0)

    add_major(doc, "8. Prioritized remediation and retest plan")
    action_rows = [
        ["1", "P1", "Add AuditLog entries and transaction boundaries to all question mutations", "Backend", "Audit cases pass; failed bulk/reorder leaves no partial state"],
        ["2", "P1", "Prevent deletion of the active administrator; require a second admin or break-glass path", "Backend", "Self-delete returns 409/403 with stable error"],
        ["3", "P2", "Validate page/limit values and return controlled 400 responses", "Backend", "abc, floats, negatives, zero, and huge values covered"],
        ["4", "P2", "Fail reorder when any ID is missing/duplicate or orders conflict", "Backend", "Atomic 400/404; no row changes on invalid batch"],
        ["5", "P2", "Stop swallowing admin API errors; render explicit retry/error states", "Frontend", "5xx/network failures cannot appear as zero data"],
        ["6", "P2", "Preserve full redirect target for customer and admin deep links", "Frontend", "All NAV-002..006 and admin subroutes return correctly"],
        ["7", "P2", "Pin OpenAI/httpx compatibility and improve personality diagnostics", "Platform", "Constructor smoke test and personality integration test pass"],
        ["8", "P2", "Repair Django/Celery test isolation and SQLite settings", "QA/Backend", "All tests pass on PostgreSQL; optional SQLite smoke works"],
        ["9", "P2", "Recalculate workbook summary and require linked evidence per automated run", "QA", "Row totals equal summary and every automated Pass has artifact/run ID"],
        ["10", "Deferred", "Retest Stripe with fresh CLI listener/webhook secret", "Payments", "Checkout/webhook/refund matrix passes in test mode"],
    ]
    add_table(doc, ["Order", "Priority", "Action", "Owner", "Exit criterion"], action_rows, [700, 1000, 3500, 1300, 2860], status_col=1, font_size=8.5)

    add_major(doc, "9. Release criteria")
    criteria = [
        ["Administrator authorization", "Pass", "Anonymous 401, non-admin 403, admin 200; UI pages render after confirmation"],
        ["Administrator auditability", "Fail", "All user and question mutations create attributable immutable audit records"],
        ["Input robustness", "Fail", "Malformed filters/pagination never generate 500"],
        ["Question data integrity", "Fail", "Reorder/bulk operations are atomic and reject stale/duplicate IDs"],
        ["Customer critical paths", "Open", "Signup, login, profile email, deep-link, processing, and chat failures retested"],
        ["Payment readiness", "Deferred", "Fresh Stripe test webhook and checkout suite passes"],
        ["Automated regression", "Fail", "All Django tests pass; browser/API artifacts retained"],
        ["Evidence integrity", "Fail", "Workbook summary reconciles exactly and unsupported passes are removed"],
    ]
    add_table(doc, ["Gate", "Current", "Required exit condition"], criteria, [2700, 1300, 5360], status_col=1, font_size=9.1)
    add_callout(doc, "Decision rule", "Release only after all P1 findings are fixed, all P0/P1 customer failures are rerun with reproducible evidence, and no critical gate remains Fail. Stripe may remain a separately tracked pre-production gate only if payment is disabled in the release candidate.", "risk")

    add_major(doc, "Appendix A. Evidence register")
    evidence_rows = [
        ["E-01", "VoiceVault_Manual_QA_Test_Cases (2).xlsx", "32 sheets / 538 cases", "Source workbook; claims audited, not accepted at face value"],
        ["E-02", "live_admin_api_results_2026-08-23.json", "49 live cases", "Machine-readable API results and timestamps"],
        ["E-03", "Hostinger runtime", "SSH, Compose, container health, Django shell", "Admin creation, health, cleanup, model checks"],
        ["E-04", "Repository f59844354a82", "Django + Next.js source", "Auth, admin, question, processing, payment, and test review"],
        ["E-05", "Local command output", "npm build/type-check; manage.py test", "Engineering gate evidence"],
        ["E-06", "In-app browser run", "Six admin routes", "Rendered UI, controls, navigation, and console evidence" if ui else "Pending action-time credential confirmation"],
    ]
    add_table(doc, ["ID", "Evidence", "Coverage", "Use"], evidence_rows, [1100, 3100, 2100, 3060], font_size=9.0)

    add_major(doc, "Appendix B. Interpretation rules")
    rules = [
        ("Pass", "Observed result directly satisfied the expected result with reproducible evidence."),
        ("Fail", "Observed result contradicted the expectation or exposed a material implementation/evidence defect."),
        ("Blocked", "A prerequisite was absent, so the intended behavior could not be evaluated."),
        ("Not Executed", "The case was intentionally not run because it was production-wide, costly, destructive, unavailable, or explicitly deferred."),
        ("Code-review finding", "The behavior is evident in source but was not necessarily triggered live when doing so would risk production data."),
        ("Inference", "A reasoned conclusion from evidence; explicitly identified and not presented as certain fact."),
    ]
    for label, value in rules:
        add_label_paragraph(doc, label, value, after=5)

    QA_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "pages_expected": "approximately 25-35", "ui_included": bool(ui), "live_api_summary": api["summary"]}, indent=2))


if __name__ == "__main__":
    build_report()
