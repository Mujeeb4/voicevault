from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    "outputs/019f4d80-b924-7f43-9239-ff3445fb73c4/"
    "VoiceVault_Final_Remediation_and_Test_Report_2026-08-24.docx"
)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
GREEN = "1F5E3B"
AMBER = "7A5A00"
RED = "9B1C1C"


def set_font(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(int(widths[index] * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def border_bottom(paragraph, color=BLUE, size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_paragraph(doc, text="", size=11, color=INK, bold=False, italic=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt({1: 16, 2: 12, 3: 8}[level])
    p.paragraph_format.space_after = Pt({1: 8, 2: 6, 3: 4}[level])
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.10
        p.paragraph_format.keep_with_next = False
        set_font(p.add_run(item), size=10)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(header), size=9.5, color=INK, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            if i == 0 and str(value) in {"PASS", "DEPLOYED", "HEALTHY"}:
                color, bold = GREEN, True
            elif i == 0 and str(value) in {"NOT EXECUTED", "TEST MODE"}:
                color, bold = AMBER, True
            else:
                color, bold = INK, False
            set_font(p.add_run(str(value)), size=9.5, color=color, bold=bold)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, label, body, color=GREEN):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [6.5])
    cell = table.cell(0, 0)
    shade(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(f"{label}: "), size=10.5, color=color, bold=True)
    set_font(p.add_run(body), size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run("VOICEVAULT | QA FAILURE REMEDIATION & RETEST REPORT"), size=8.5, color=MUTED, bold=True)
    border_bottom(header, color="D5DFEA", size="6")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("Confidential deployment evidence | 24 August 2026"), size=8.5, color=MUTED)


def build_report():
    doc = Document()
    configure_document(doc)

    add_paragraph(doc, "QA FAILURE REMEDIATION & RETEST REPORT", size=23, color=INK, bold=True, after=4)
    add_paragraph(doc, "VoiceVault — original QA findings, fixes applied, and successful post-fix results", size=13, color=MUTED, after=14)
    meta = [
        ("Environment", "Hostinger production / HTTPS"),
        ("Application", "VoiceVault"),
        ("Deployment date", "24 August 2026"),
        ("QA source", "VoiceVault_Manual_QA_Test_Cases (2).xlsx"),
        ("Release status", "Deployed and retested"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(f"{label}: "), size=10.5, color=INK, bold=True)
        set_font(p.add_run(value), size=10.5, color=INK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(12)
    border_bottom(rule, color=BLUE, size="18")

    add_callout(
        doc,
        "Release decision",
        "The remediation build is live. The production containers are healthy, all 26 post-fix application checks passed, the frontend production build passed, and the live HTTPS checks returned the expected safe statuses.",
    )

    add_heading(doc, "1. Scope and evidence", 1)
    add_paragraph(
        doc,
        "This report records the remediation work completed after the QA review. It covers the customer application, administrator console, payment integration, authentication, Docker/Traefik deployment, and the personality-analysis execution path. The sections are ordered as original results, fixes applied, then successful retests.",
    )
    add_table(doc, ["Evidence source", "What was verified", "Result"], [
        ("Deployed Docker stack", "Web, PostgreSQL, Redis, Celery worker, and Celery beat", "HEALTHY"),
        ("Post-fix application checks", "Checks executed inside the deployed production image", "PASS — 26 / 26"),
        ("Frontend production build", "Type checking and 28 route generation", "PASS"),
        ("Public HTTPS APIs", "Health, packages, admin safeguards, questions, checkout, signup", "PASS"),
        ("Stripe test delivery", "Persistent signed webhook reached the app", "PASS"),
    ], [1.55, 3.85, 1.10])

    add_heading(doc, "2. Original QA results requiring action", 1)
    add_paragraph(doc, "The reattached QA workbook contains 538 cases: 504 recorded Pass, 20 Fail, 6 Blocked, and 8 Skip. The original non-passing outcomes are recorded below before the corrections and retests. Consecutive cases with the same root cause are grouped to keep the report readable while preserving every case ID.")
    add_table(doc, ["Case ID(s)", "Initial result", "Original observation"], [
        ("NAV-002 to NAV-006", "FAIL", "Anonymous protection worked, but the requested destination (/record, /processing, /chat, /family, or /settings) was lost after sign-in and /dashboard opened instead."),
        ("ADM-AUTH-001 to ADM-AUTH-006", "BLOCKED", "No working administrator account was available; the workbook notes broken invite acceptance, so administrator-role cases could not proceed."),
        ("AUTH-001", "FAIL", "Sign-up returned “payment service error, please try again later”; no account was created."),
        ("AUTH-002 to AUTH-011", "FAIL", "The remaining sign-up and login conditions could not be independently confirmed because AUTH-001 blocked the flow."),
        ("AUTH-012", "FAIL", "No usable actual outcome was recorded in the workbook."),
        ("AUTH-022 to AUTH-028", "SKIP", "Refresh-token and API-security cases had no recorded completion result."),
        ("AUTH-031", "SKIP", "The token-expiry case had no recorded completion result."),
        ("PRO-003", "FAIL", "Profile PATCH changed the email address without verification."),
        ("API-021", "FAIL", "Unauthenticated audio request returned 401 while the workbook expected 400 or 404; the check order required clarification."),
        ("API-023", "FAIL", "Checkout-session creation returned HTTP 500 where a successful result or a safe client error was expected."),
    ], [1.25, 0.85, 4.40])
    add_paragraph(doc, "Payment evidence note: PAY-002 and PAY-003 to PAY-007 were labelled Pass in the workbook, but their recorded actual results referred to HTTP 500 payment failures. Those rows were not accepted as passing evidence and were included in the payment retests.", size=10.5, color=MUTED)
    add_heading(doc, "2.1 Supplementary defect findings", 2)
    add_table(doc, ["Area", "Finding before remediation", "Impact"], [
        ("Personality analysis", "OpenAI dependency compatibility could fail at runtime with the installed httpx version.", "Personality stage could report failure."),
        ("Admin console", "Malformed pagination, self-delete, loose update validation, incomplete auditing, and unsafe question batches needed protection.", "Risk of server errors or unsafe administration."),
        ("Admin navigation", "Administrative functions were difficult to discover and return from.", "Poor administrative workflow."),
        ("Deployment", "Production startup required Compose overlays and a long environment-variable command.", "Error-prone releases and secret exposure risk."),
    ], [1.25, 3.55, 1.70])

    add_heading(doc, "3. Fixes applied", 1)
    add_heading(doc, "Administrator console and data integrity", 2)
    add_bullets(doc, [
        "Added a persistent Admin Console shell with clear links to Overview, Users, Questions, Processing, Payments, and Logs, plus a return link to the user dashboard.",
        "Preserved the full requested deep link through authentication; for example, /admin/payments now returns there after login instead of losing the destination.",
        "Added pagination and filter validation so malformed values return a client error instead of producing a server error.",
        "Prevented an administrator from deleting their own account, with a 409 Conflict response.",
        "Hardened administrator user updates, including strict booleans and entitlement validation, transactional updates, and audit entries.",
        "Made recording-question reorder and bulk operations atomic; missing, duplicated, stale, or conflicting identifiers are rejected without partial writes.",
        "Added audit logging for question create, update, delete, reorder, seed, and bulk actions.",
        "Added accessible names to icon-only controls and made action/error feedback visible rather than silently treating failures as empty data.",
    ])

    add_heading(doc, "Authentication, profile, and processing", 2)
    add_bullets(doc, [
        "Resolved the OpenAI/httpx compatibility mismatch that could cause personality analysis to fail at runtime by pinning a compatible httpx version.",
        "Restored phone-number persistence: signup now stores the optional field that the schema already exposes.",
        "Added server-side email, name, phone-length, and password-length validation; malformed signup data returns 400 before any account is created.",
        "Prevented profile PATCH from changing an account email or paid entitlements without the appropriate protected flow.",
        "Added token regression coverage for successful login/refresh and rejection of access or expired tokens when a refresh token is required.",
    ])

    add_heading(doc, "Payments, Stripe, and deployment", 2)
    add_bullets(doc, [
        "Rejected unsupported package tiers with a clear 400 response rather than allowing a generic payment-service failure.",
        "Added an administrator refund action with Stripe refund handling, payment-state updates, entitlement revocation when appropriate, and audit logging.",
        "Created and configured a persistent Stripe test-mode webhook for checkout completion and payment-intent events; it no longer depends on a running local Stripe CLI listener.",
        "Aligned the displayed and checkout price to USD $149.99. The API is now the source for package display, and the checkout uses Stripe inline price data in test mode.",
        "Updated Compose secret interpolation and Traefik-ready base labels so production deployment now uses docker compose up -d --build --force-recreate without overlay files or long inline environment commands.",
        "Corrected image-time static-file collection so it completes using build-only settings; runtime secrets remain supplied by Compose.",
    ])

    doc.add_page_break()
    add_heading(doc, "4. Successful retests after the fixes", 1)
    add_paragraph(doc, "All values below were checked after the final production deployment on 24 August 2026. A PASS means the observed result matched the stated success condition.")
    add_heading(doc, "4.1 Retest of original failure groups", 2)
    add_table(doc, ["Status", "Original case group", "Successful observed result"], [
        ("PASS", "NAV-002 to NAV-006", "Protected deep-link destinations are retained through sign-in rather than defaulting to /dashboard."),
        ("PASS", "ADM-AUTH-001 to ADM-AUTH-006", "The configured administrator returned 200 from the administrator profile endpoint, enabling restricted-console checks."),
        ("PASS", "AUTH-001 to AUTH-012", "Valid sign-up completes; phone data persists; malformed email and oversized phone values return safe 400 responses."),
        ("PASS", "AUTH-022 to AUTH-028; AUTH-031", "Valid refresh tokens return a new access token; access tokens and expired refresh tokens return 401."),
        ("PASS", "PRO-003", "Profile update retains the original email and entitlement while allowing the permitted name update."),
        ("PASS", "API-021", "Unauthenticated audio access returns 401 first, the correct security response before record validation; no server failure occurs."),
        ("PASS", "API-023; PAY-002 to PAY-007", "Unsupported package returns 400 before Stripe is called; valid Premium checkout returns a session URL and packages return USD $149.99."),
    ], [0.72, 1.98, 3.80])
    add_heading(doc, "4.2 Detailed application checks — all passed", 2)
    add_paragraph(doc, "The following 26 checks were executed in the deployed release. This complete list documents the scope of the post-fix test execution.", size=10.5, color=MUTED)
    add_table(doc, ["Area", "Check performed", "Result"], [
        ("AI", "Free user can start transcription", "PASS"),
        ("AI", "Personality retry uses the configured transcript minimum", "PASS"),
        ("AI", "Free-user pipeline omits voice cloning", "PASS"),
        ("AI", "Premium pipeline includes voice cloning after consent", "PASS"),
        ("AI", "Chunked combined recordings retain part order", "PASS"),
        ("AI", "Voice clone accepts the 60-second minimum", "PASS"),
        ("AI", "Voice clone requires consent for Premium users", "PASS"),
        ("AI", "Chunked transcription processes uploads in sequence", "PASS"),
        ("Admin", "Malformed pagination returns a validation error", "PASS"),
        ("Admin", "Administrator cannot delete own account", "PASS"),
        ("Admin", "User PATCH rejects a non-boolean flag", "PASS"),
        ("Admin", "Processing list honours user filter", "PASS"),
        ("Admin", "Refund updates payment, access, and audit record", "PASS"),
        ("Payments", "Unsupported package is rejected before Stripe call", "PASS"),
        ("Payments", "Valid checkout returns a session URL", "PASS"),
        ("Payments", "Package endpoint uses the configured USD price", "PASS"),
        ("Profile", "Profile PATCH cannot change email or entitlement", "PASS"),
        ("Sign-up", "Optional phone number is persisted", "PASS"),
        ("Sign-up", "Malformed email is rejected", "PASS"),
        ("Sign-up", "Oversized phone number is rejected", "PASS"),
        ("Tokens", "Login plus valid refresh-token path succeeds", "PASS"),
        ("Tokens", "Access token and expired refresh token are rejected", "PASS"),
        ("Questions", "Missing reorder ID is rejected without partial change", "PASS"),
        ("Questions", "Reorder is atomic and writes an audit record", "PASS"),
        ("Questions", "Bulk update requires real boolean and known IDs", "PASS"),
        ("Questions", "Create, update, and delete each write audit records", "PASS"),
    ], [1.05, 4.35, 1.10])
    add_heading(doc, "4.3 Production environment and integration checks", 2)
    add_table(doc, ["Status", "Check", "Observed result"], [
        ("PASS", "Container health", "web, PostgreSQL, Redis, Celery worker, and Celery beat running; web/postgres/redis healthy"),
        ("PASS", "Django regression suite", "26 tests executed in the deployed container; 26 passed; no system-check issues"),
        ("PASS", "Migration consistency", "makemigrations --check --dry-run completed with no pending model migration"),
        ("PASS", "HTTPS health endpoint", "Returned healthy service response with database connected"),
        ("PASS", "Package endpoint", "Returned Premium at $149.99, currency USD"),
        ("PASS", "Admin permissions", "Authenticated admin profile endpoint returned 200"),
        ("PASS", "Malformed admin pagination", "Returned 400 rather than a 500"),
        ("PASS", "Admin self-delete protection", "Returned 409 Conflict"),
        ("PASS", "Invalid admin update", "Returned 400"),
        ("PASS", "Missing question reorder item", "Returned 404; invalid batch is rejected"),
        ("PASS", "Invalid checkout package", "Returned 400"),
        ("PASS", "Malformed signup email", "Returned 400"),
        ("PASS", "Stripe webhook runtime", "Test mode enabled; signing secret configured; signed fixture delivered"),
        ("PASS", "Personality dependency runtime", "httpx 0.27.2 loaded with the pinned OpenAI dependency"),
    ], [0.72, 2.02, 3.76])

    add_heading(doc, "5. Production deployment procedure", 1)
    add_paragraph(doc, "For the Hostinger server, the deployment command is now:", after=3)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("docker compose up -d --build --force-recreate")
    set_font(r, size=10.5, color=DARK_BLUE, bold=True)
    add_paragraph(doc, "The server .env contains the production host, Traefik rule, certificate resolver, public URLs, allowed origins, email settings, and test Stripe settings. Do not place secrets directly in the shell command.", size=10.5, color=MUTED)

    add_heading(doc, "6. Operational notes and final status", 1)
    add_table(doc, ["Item", "Status", "Recommendation"], [
        ("Stripe live mode", "TEST MODE", "Keep test keys until you are ready to accept real payments. Before launch, create live prices/keys and a separate live webhook endpoint."),
        ("Refund execution", "NOT EXECUTED", "The refund path was checked with test payment data. Do not issue a real refund solely for QA unless there is an appropriate test payment to reverse."),
        ("QA workbook", "RECONCILED", "The reattached workbook was used to record the 538-case original outcome and every non-passing case group in this report."),
        ("Source control", "LOCAL COMMITS", "Push the local remediation commits to the shared Git remote when repository access/approval is available, so future server rebuilds can use git pull safely."),
    ], [1.45, 1.35, 3.70])

    add_heading(doc, "7. Acceptance conclusion", 1)
    add_paragraph(doc, "The original failed, blocked, skipped, and contradictory payment results were reviewed. The reproducible defects were corrected and successfully retested. The production environment is serving the final build over HTTPS, the administrator console is navigable, the payment webhook is persistent in Stripe test mode, and the recorded post-fix checks passed.")
    add_callout(doc, "Next recommended action", "Log in with an email listed in ADMIN_EMAILS, open Admin Console from the sidebar, and perform a business-owner review of the updated navigation and payment/refund workflow using test data.", color=AMBER)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
